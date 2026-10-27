from datetime import timedelta
from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.db.models import Avg, Sum, Max, Min, Count
from django.shortcuts import get_object_or_404, redirect, render
from django.http import HttpResponse
from django.utils import timezone

from students.models import Student
from accounts.models import SchoolBranding
from accounts.staff_access import (
    is_admin_user,
    get_staff_class_pairs,
    get_staff_students,
    can_view_class,
    can_enter_marks,
    get_staff_subjects_for_class,
    is_class_teacher_for,
)
from fees.models import AcademicYear, Term

from .models import (
    Stream,
    Subject,
    AssessmentType,
    Assessment,
    AssessmentWindow,
)

from .performance_analytics import build_performance_analytics

from .mark_permissions import (
    is_admin,
    can_enter_marks,
    can_access_class,
    allowed_subject_ids_for_class,
)



# ============================================================
# HELPER FUNCTIONS
# ============================================================

def get_class_names():
    """
    Get all classes currently represented in the Student table.
    """
    return (
        Student.objects
        .exclude(class_name__isnull=True)
        .exclude(class_name__exact="")
        .values_list("class_name", flat=True)
        .distinct()
        .order_by("class_name")
    )


def get_stream_names():
    """
    Get active streams.
    """
    return Stream.objects.filter(
        is_active=True
    ).order_by("name")


def get_staff_class_names(user):
    """
    Return classes available to the current user.
    Admins see all classes.
    Teachers see only classes containing students
    they are assigned to.
    """
    if is_admin_user(user):
        return get_class_names()

    return (
        get_staff_students(user)
        .exclude(class_name__isnull=True)
        .exclude(class_name__exact="")
        .values_list("class_name", flat=True)
        .distinct()
        .order_by("class_name")
    )


def get_staff_stream_names(user, class_name=""):
    """
    Return streams available to the current user.
    Admins see all active streams.
    Teachers see only streams containing students
    within their assigned classes/streams.
    """
    if is_admin_user(user):
        return get_stream_names()

    students = get_staff_students(user)

    if class_name:
        students = students.filter(class_name=class_name)

    streams = (
        students
        .exclude(stream__isnull=True)
        .exclude(stream__exact="")
        .values_list("stream", flat=True)
        .distinct()
        .order_by("stream")
    )

    return streams

def get_students_for_class(class_name="", stream_name=""):
    """
    Return learners belonging to a selected class
    and optionally selected stream.
    """
    queryset = Student.objects.all()

    if class_name:
        queryset = queryset.filter(
            class_name=class_name
        )

    if stream_name:
        queryset = queryset.filter(
            stream=stream_name
        )

    return queryset.order_by(
        "first_name",
        "middle_name",
        "last_name",
    )


def get_student_stream_name(student):
    """
    Safely get the learner's stream.

    Supports projects where Student.stream is either:
    - a text field
    - a ForeignKey to Stream
    """
    stream = getattr(student, "stream", "")

    if not stream:
        return ""

    if hasattr(stream, "name"):
        return stream.name

    return str(stream)



# ============================================================
# MARK PERMISSION HELPERS
# ============================================================

def _mark_student_stream(student):
    value = getattr(student, "stream", "")
    if hasattr(value, "name"):
        return value.name or ""
    return value or ""


def _authorized_subjects_for_class(user, class_name, stream=""):
    """
    Return the subjects this user may enter for a class/stream.
    None means unrestricted.
    """
    allowed_ids = allowed_subject_ids_for_class(
        user,
        class_name,
        stream,
    )

    if allowed_ids is None:
        return None

    return allowed_ids


# ============================================================
# ACADEMIC DASHBOARD
# ============================================================

@login_required
def academic_dashboard(request):

    from accounts.views import user_has_role_permission

    if not (
        is_admin_user(request.user)
        or user_has_role_permission(
            request.user,
            "view_academicrecord",
            "academic",
        )
    ):
        return redirect("accounts:staff_dashboard")
    # --------------------------------------------------------
    # STAFF: open the staff academic class directly
    # --------------------------------------------------------
    if (
        not is_admin_user(request.user)
        and getattr(
            getattr(request.user, "profile", None),
            "role",
            ""
        ) == "STAFF"
    ):
        class_pairs = get_staff_class_pairs(request.user)

        if len(class_pairs) == 1:
            class_name, stream = class_pairs[0]

            from django.urls import reverse
            from urllib.parse import urlencode

            url = reverse(
                "accounts:staff_academic_class"
            )

            query = urlencode({
                "class_name": class_name,
                "stream": stream,
            })

            return redirect(f"{url}?{query}")

        if class_pairs:
            return redirect(
                "accounts:staff_academic_portal"
            )

        return redirect(
            "accounts:staff_dashboard"
        )

    # --------------------------------------------------------
    # ADMIN: retain the normal Academic Dashboard
    # --------------------------------------------------------

    students_count = Student.objects.count()

    subjects_count = Subject.objects.filter(
        is_active=True
    ).count()

    assessment_types_count = AssessmentType.objects.filter(
        is_active=True
    ).count()

    results_count = Assessment.objects.count()

    streams_count = Stream.objects.filter(
        is_active=True
    ).count()

    return render(
        request,
        "academic/dashboard.html",
        {
            "students_count": students_count,
            "subjects_count": subjects_count,
            "assessment_types_count": assessment_types_count,
            "results_count": results_count,
            "streams_count": streams_count,
        }
    )


# ============================================================
# MARKS ENTRY
# ============================================================

# MARKS ENTRY
# ============================================================

@login_required
def marks_entry(request):

    # ========================================================
    # BASE DATA
    # ========================================================

    # ========================================================
    # ASSESSMENT CONTROL CENTRE
    # Only currently recordable assessment windows are exposed
    # to Mark Entry.
    #
    # OPEN          -> allowed
    # CLOSING SOON  -> allowed
    # REOPENED      -> allowed
    # CLOSED        -> hidden
    # DEADLINE      -> hidden
    # ========================================================

    from .models import AssessmentWindow

    open_windows = (
        AssessmentWindow.objects
        .filter(is_active=True)
        .select_related(
            "academic_year",
            "term",
            "assessment_type",
        )
        .order_by(
            "-deadline"
        )
    )

    # Only windows whose recording is currently allowed.
    open_windows = [
        window
        for window in open_windows
        if window.recording_open()
    ]

    # Remove duplicates while preserving order.
    available_years = []
    available_terms = []
    available_assessment_types = []

    seen_years = set()
    seen_terms = set()
    seen_assessments = set()

    for window in open_windows:

        if window.academic_year_id not in seen_years:
            available_years.append(window.academic_year)
            seen_years.add(window.academic_year_id)

        if window.term_id not in seen_terms:
            available_terms.append(window.term)
            seen_terms.add(window.term_id)

        if window.assessment_type_id not in seen_assessments:
            available_assessment_types.append(
                window.assessment_type
            )
            seen_assessments.add(
                window.assessment_type_id
            )

    # Backwards-compatible template variables.
    # The Mark Entry template can continue using these names.
    assessment_types = available_assessment_types
    academic_years = available_years
    terms = available_terms

    # --------------------------------------------------------
    # STAFF-RESTRICTED CLASSES / STUDENTS
    # --------------------------------------------------------

    allowed_students = get_staff_students(request.user)

    class_names = (
        allowed_students
        .exclude(class_name__isnull=True)
        .exclude(class_name__exact="")
        .values_list("class_name", flat=True)
        .distinct()
        .order_by("class_name")
    )

    streams = (
        allowed_students
        .values_list("stream", flat=True)
        .distinct()
        .order_by("stream")
    )

    # --------------------------------------------------------
    # Subjects will be determined by selected class.
    # Admin initially sees all subjects.
    # --------------------------------------------------------

    if is_admin_user(request.user):

        subjects = Subject.objects.filter(
            is_active=True
        ).order_by("name")

    else:

        subjects = Subject.objects.none()

    students = allowed_students.order_by(
        "class_name",
        "first_name",
        "middle_name",
        "last_name",
    )

    selected_student = None
    selected_year = None
    selected_term = None
    selected_assessment_type = None
    selected_subject = None

    selected_class = ""
    selected_stream = ""

    selected_subjects = []

    individual_assessments = []
    class_students = []
    class_subject_marks = []
    all_subject_marks = []

    all_subjects_selected = False

    mode = request.GET.get(
        "mode",
        "select"
    )

    # ========================================================
    # VALIDATE SELECTED ASSESSMENT WINDOW
    # ========================================================
    #
    # A Year + Term + Assessment Type combination is valid only
    # when the exact combination has an active recordable window.
    # ========================================================

    requested_year_id = request.GET.get("year")
    requested_term_id = request.GET.get("term")
    requested_assessment_type_id = request.GET.get(
        "assessment_type"
    )

    if (
        requested_year_id
        and requested_term_id
        and requested_assessment_type_id
    ):

        valid_window = next(
            (
                window
                for window in open_windows
                if str(window.academic_year_id)
                == str(requested_year_id)
                and str(window.term_id)
                == str(requested_term_id)
                and str(window.assessment_type_id)
                == str(requested_assessment_type_id)
            ),
            None,
        )

        if not valid_window:

            messages.error(
                request,
                "Mark entry is unavailable. The selected "
                "assessment is not currently open for recording."
            )

            return redirect(
                "academic:marks_entry"
            )

    # ========================================================
    # INDIVIDUAL MARKS
    # ========================================================

    if mode == "individual":

        student_id = request.GET.get("student")
        year_id = request.GET.get("year")
        term_id = request.GET.get("term")
        assessment_type_id = request.GET.get(
            "assessment_type"
        )

        if student_id:

            selected_student = get_object_or_404(
                Student,
                id=student_id
            )

            # ------------------------------------------------
            # SECURITY:
            # Teacher must have access to this learner.
            # ------------------------------------------------

            if not is_admin_user(request.user):

                if not allowed_students.filter(
                    pk=selected_student.pk
                ).exists():

                    messages.error(
                        request,
                        "You are not authorized to enter "
                        "marks for this learner."
                    )

                    return redirect(
                        "academic:marks_entry"
                    )

        if year_id:

            selected_year = get_object_or_404(
                AcademicYear,
                id=year_id
            )

        if term_id:

            selected_term = get_object_or_404(
                Term,
                id=term_id
            )

        if assessment_type_id:

            selected_assessment_type = get_object_or_404(
                AssessmentType,
                id=assessment_type_id
            )

        # ----------------------------------------------------
        # Determine class/stream from learner.
        # ----------------------------------------------------

        if selected_student:

            selected_class = (
                selected_student.class_name or ""
            )

            selected_stream = get_student_stream_name(
                selected_student
            )

            if not can_view_class(
                request.user,
                selected_class,
                selected_stream,
            ):

                messages.error(
                    request,
                    "You are not authorized to access "
                    "this class."
                )

                return redirect(
                    "academic:marks_entry"
                )

            subjects = get_staff_subjects_for_class(
                request.user,
                selected_class,
                selected_stream,
            )

        # ----------------------------------------------------
        # LOAD EXISTING MARKS
        # ----------------------------------------------------


        if selected_student:
            student_stream = _mark_student_stream(selected_student)

            if not can_access_class(
                request.user,
                selected_student.class_name,
                student_stream,
            ):
                messages.error(
                    request,
                    "You are not authorized to enter marks for this learner's class.",
                )
                return redirect("academic:marks_entry")

        if (
            selected_student
            and selected_year
            and selected_term
            and selected_assessment_type
        ):

            # MARK INDIVIDUAL PERMISSION CHECK

            allowed_ids = _authorized_subjects_for_class(
                request.user,
                selected_student.class_name,
                _mark_student_stream(selected_student),
            )

            if allowed_ids is not None:
                subjects_for_entry = subjects.filter(
                    id__in=allowed_ids
                )
            else:
                subjects_for_entry = subjects

            # MARK SUBJECT PERMISSION FILTER
            for subject in subjects_for_entry:

                assessment = Assessment.objects.filter(
                    student=selected_student,
                    academic_year=selected_year,
                    term=selected_term,
                    subject=subject,
                    assessment_type=selected_assessment_type,
                ).first()

                individual_assessments.append({
                    "subject": subject,
                    "assessment": assessment,
                })


    # ========================================================
    # WHOLE CLASS MARKS
    # ========================================================

    elif mode == "class":

        year_id = request.GET.get("year")
        term_id = request.GET.get("term")
        assessment_type_id = request.GET.get(
            "assessment_type"
        )

        selected_class = request.GET.get(
            "class_name",
            ""
        ).strip()

        selected_stream = request.GET.get(
            "stream",
            ""
        ).strip()

        subject_id = request.GET.get(
            "subject",
            ""
        )

        if year_id:

            selected_year = get_object_or_404(
                AcademicYear,
                id=year_id
            )

        if term_id:

            selected_term = get_object_or_404(
                Term,
                id=term_id
            )

        if assessment_type_id:

            selected_assessment_type = get_object_or_404(
                AssessmentType,
                id=assessment_type_id
            )

        # ----------------------------------------------------
        # CLASS SECURITY
        # ----------------------------------------------------

        if selected_class:

            if not can_view_class(
                request.user,
                selected_class,
                selected_stream,
            ):

                messages.error(
                    request,
                    "You are not authorized to access "
                    "this class."
                )

                return redirect(
                    "academic:marks_entry"
                )

            subjects = get_staff_subjects_for_class(
                request.user,
                selected_class,
                selected_stream,
            )

        # ----------------------------------------------------
        # SUBJECT
        # ----------------------------------------------------

        if subject_id == "all":

            # Only admins and class teachers can use ALL.
            if not is_admin_user(request.user):

                if not get_staff_subjects_for_class(
                    request.user,
                    selected_class,
                    selected_stream,
                ).count():

                    messages.error(
                        request,
                        "You are not authorized to enter "
                        "marks for this class."
                    )

                    return redirect(
                        "academic:marks_entry"
                    )

                # A subject teacher must never receive
                # the ALL option.

                if not is_class_teacher_for(
                    request.user,
                    selected_class,
                    selected_stream,
                ):

                    messages.error(
                        request,
                        "Subject teachers can only enter "
                        "their assigned subjects."
                    )

                    return redirect(
                        "academic:marks_entry"
                    )

            selected_subjects = list(
                subjects
            )

            selected_subject = None
            all_subjects_selected = True

        elif subject_id:

            selected_subject = get_object_or_404(
                Subject,
                id=subject_id,
                is_active=True
            )

            # ------------------------------------------------
            # SUBJECT SECURITY
            # ------------------------------------------------

            if not can_enter_marks(
                request.user,
                selected_class,
                selected_stream,
                selected_subject,
            ):

                messages.error(
                    request,
                    "You are not authorized to enter "
                    "marks for this subject and class."
                )

                return redirect(
                    "academic:marks_entry"
                )

            selected_subjects = [
                selected_subject
            ]

        # ----------------------------------------------------
        # LOAD CLASS MARKS
        # ----------------------------------------------------

        if (
            selected_year
            and selected_term
            and selected_assessment_type
            and selected_class
            and selected_subjects
        ):

            class_students = list(
                allowed_students.filter(
                    class_name=selected_class,
                    stream=selected_stream,
                ).order_by(
                    "first_name",
                    "middle_name",
                    "last_name",
                )
            )

            # =================================================
            # ALL SUBJECTS
            # =================================================

            if subject_id == "all":

                for student in class_students:

                    student_subject_marks = []

                    for subject in selected_subjects:

                        assessment = Assessment.objects.filter(
                            student=student,
                            academic_year=selected_year,
                            term=selected_term,
                            subject=subject,
                            assessment_type=selected_assessment_type,
                        ).first()

                        student_subject_marks.append({
                            "subject": subject,
                            "assessment": assessment,
                        })

                    all_subject_marks.append({
                        "student": student,
                        "subjects": student_subject_marks,
                    })

                    class_subject_marks.append({
                        "student": student,
                        "subjects": student_subject_marks,
                    })

            # =================================================
            # SINGLE SUBJECT
            # =================================================

            else:

                for student in class_students:

                    assessment = Assessment.objects.filter(
                        student=student,
                        academic_year=selected_year,
                        term=selected_term,
                        subject=selected_subject,
                        assessment_type=selected_assessment_type,
                    ).first()

                    class_subject_marks.append({
                        "student": student,
                        "assessment": assessment,
                    })


    # ========================================================
    # SAVE MARKS
    # ========================================================

    if request.method == "POST":

        # ====================================================
        # ASSESSMENT WINDOW ENFORCEMENT
        # Only OPEN, CLOSING SOON and REOPENED assessments
        # may receive marks. This applies to ALL mark-entry
        # methods, including individual and whole-class entry.
        # ====================================================

        from .models import AssessmentWindow

        window = AssessmentWindow.objects.filter(
            academic_year_id=request.POST.get("year"),
            term_id=request.POST.get("term"),
            assessment_type_id=request.POST.get("assessment_type"),
            is_active=True,
        ).first()

        if not window:
            messages.error(
                request,
                "Mark entry is blocked. No active assessment "
                "window has been opened for this assessment."
            )
            return redirect("academic:marks_entry")

        if not window.recording_open():
            if window.is_reopened:
                window_status_message = (
                    "Mark entry is currently unavailable."
                )
            else:

                if timezone.now() >= window.deadline:
                    window_status_message = (
                        "This assessment deadline has passed. "
                        "Mark entry is closed."
                    )
                else:
                    window_status_message = (
                        "This assessment is currently closed. "
                        "Mark entry is not allowed."
                    )

            messages.error(
                request,
                window_status_message
            )
            return redirect("academic:marks_entry")

        post_mode = request.POST.get(
            "mode"
        )

        year_id = request.POST.get(
            "year"
        )

        term_id = request.POST.get(
            "term"
        )

        assessment_type_id = request.POST.get(
            "assessment_type"
        )

        academic_year = get_object_or_404(
            AcademicYear,
            id=year_id
        )

        term = get_object_or_404(
            Term,
            id=term_id
        )

        assessment_type = get_object_or_404(
            AssessmentType,
            id=assessment_type_id
        )

        # ====================================================
        # SAVE INDIVIDUAL
        # ====================================================

        if post_mode == "individual":

            student_id = request.POST.get(
                "student"
            )

            student = get_object_or_404(
                Student,
                id=student_id
            )

            # ------------------------------------------------
            # SECURITY
            # ------------------------------------------------

            if not is_admin_user(request.user):

                if not allowed_students.filter(
                    pk=student.pk
                ).exists():

                    messages.error(
                        request,
                        "You are not authorized to enter "
                        "marks for this learner."
                    )

                    return redirect(
                        "academic:marks_entry"
                    )

            student_class = (
                student.class_name or ""
            )

            student_stream = get_student_stream_name(
                student
            )

            # ------------------------------------------------
            # SAVE ONLY AUTHORIZED SUBJECTS
            # ------------------------------------------------

            allowed_subjects = get_staff_subjects_for_class(
                request.user,
                student_class,
                student_stream,
            )

            for subject in allowed_subjects:

                score_value = request.POST.get(
                    f"score_{subject.id}"
                )

                comment = request.POST.get(
                    f"comment_{subject.id}",
                    ""
                )

                if score_value in (
                    None,
                    ""
                ):
                    continue

                # Explicit backend permission check.
                if not can_enter_marks(
                    request.user,
                    student_class,
                    student_stream,
                    subject,
                ):
                    continue

                try:

                    score = float(
                        score_value
                    )

                except (
                    TypeError,
                    ValueError
                ):

                    continue

                score = max(
                    0,
                    min(score, 100)
                )

                Assessment.objects.update_or_create(
                    student=student,
                    academic_year=academic_year,
                    term=term,
                    subject=subject,
                    assessment_type=assessment_type,
                    defaults={
                        "score": score,
                        "teacher_comment": comment,
                    }
                )

            return redirect(
                f"/academic/marks/"
                f"?mode=individual"
                f"&student={student.id}"
                f"&year={academic_year.id}"
                f"&term={term.id}"
                f"&assessment_type={assessment_type.id}"
            )

        # ====================================================
        # SAVE CLASS
        # ====================================================

        elif post_mode == "class":

            selected_class = request.POST.get(
                "class_name",
                ""
            ).strip()

            selected_stream = request.POST.get(
                "stream",
                ""
            ).strip()

            subject_id = request.POST.get(
                "subject",
                ""
            ).strip()

            # ------------------------------------------------
            # CLASS SECURITY
            # ------------------------------------------------

            if not can_view_class(
                request.user,
                selected_class,
                selected_stream,
            ):

                messages.error(
                    request,
                    "You are not authorized to enter "
                    "marks for this class."
                )

                return redirect(
                    "academic:marks_entry"
                )

            class_students = (
                allowed_students
                .filter(
                    class_name=selected_class,
                    stream=selected_stream,
                )
                .order_by(
                    "first_name",
                    "middle_name",
                    "last_name",
                )
            )

            # =================================================
            # SAVE ALL SUBJECTS
            # =================================================

            if subject_id == "all":

                if not is_admin_user(request.user):

                    if not is_class_teacher_for(
                        request.user,
                        selected_class,
                        selected_stream,
                    ):

                        messages.error(
                            request,
                            "Only a class teacher or admin "
                            "can enter marks for all subjects."
                        )

                        return redirect(
                            "academic:marks_entry"
                        )

                selected_subjects = (
                    get_staff_subjects_for_class(
                        request.user,
                        selected_class,
                        selected_stream,
                    )
                )

                for student in class_students:

                    for subject in selected_subjects:

                        if not can_enter_marks(
                            request.user,
                            selected_class,
                            selected_stream,
                            subject,
                        ):
                            continue

                        score_value = request.POST.get(
                            f"score_{student.id}_{subject.id}"
                        )

                        comment = request.POST.get(
                            f"comment_{student.id}_{subject.id}",
                            ""
                        )

                        if score_value in (
                            None,
                            ""
                        ):
                            continue

                        try:

                            score = float(
                                score_value
                            )

                        except (
                            TypeError,
                            ValueError
                        ):

                            continue

                        score = max(
                            0,
                            min(score, 100)
                        )

                        Assessment.objects.update_or_create(
                            student=student,
                            academic_year=academic_year,
                            term=term,
                            subject=subject,
                            assessment_type=assessment_type,
                            defaults={
                                "score": score,
                                "teacher_comment": comment,
                            }
                        )

                messages.success(request, "Marks saved successfully.")
                return redirect(
                    f"/academic/marks/"
                    f"?mode=class"
                    f"&year={academic_year.id}"
                    f"&term={term.id}"
                    f"&assessment_type={assessment_type.id}"
                    f"&class_name={selected_class}"
                    f"&stream={selected_stream}"
                    f"&subject=all"
                )

            # =================================================
            # SAVE ONE SUBJECT
            # =================================================

            selected_subject = get_object_or_404(
                Subject,
                id=subject_id,
                is_active=True
            )

            # ------------------------------------------------
            # CRITICAL BACKEND SECURITY
            # ------------------------------------------------

            if not can_enter_marks(
                request.user,
                selected_class,
                selected_stream,
                selected_subject,
            ):

                messages.error(
                    request,
                    "You are not authorized to enter "
                    "marks for this subject and class."
                )

                return redirect(
                    "academic:marks_entry"
                )

            for student in class_students:

                score_value = request.POST.get(
                    f"score_{student.id}"
                )

                comment = request.POST.get(
                    f"comment_{student.id}",
                    ""
                )

                if score_value in (
                    None,
                    ""
                ):
                    continue

                try:

                    score = float(
                        score_value
                    )

                except (
                    TypeError,
                    ValueError
                ):

                    continue

                score = max(
                    0,
                    min(score, 100)
                )

                Assessment.objects.update_or_create(
                    student=student,
                    academic_year=academic_year,
                    term=term,
                    subject=selected_subject,
                    assessment_type=assessment_type,
                    defaults={
                        "score": score,
                        "teacher_comment": comment,
                    }
                )

            messages.success(request, "Marks saved successfully.")
            return redirect(
                f"/academic/marks/"
                f"?mode=class"
                f"&year={academic_year.id}"
                f"&term={term.id}"
                f"&assessment_type={assessment_type.id}"
                f"&class_name={selected_class}"
                f"&stream={selected_stream}"
                f"&subject={selected_subject.id}"
            )

    # ========================================================
    # RENDER
    # ========================================================


    # --------------------------------------------------------
    # ASSESSMENT MANAGEMENT / RECORDING WINDOW DATA
    # Shared with the Assessment & Marks Management centre.
    # --------------------------------------------------------

    assessment_data = []

    for window in open_windows:

        total_students = Student.objects.filter(
            class_name__isnull=False
        ).exclude(
            class_name=""
        ).count()

        recorded = Assessment.objects.filter(
            academic_year=window.academic_year,
            term=window.term,
            assessment_type=window.assessment_type,
        ).count()

        subjects_count = Subject.objects.filter(
            is_active=True
        ).count()

        expected = total_students * subjects_count

        percentage = (
            round((recorded / expected) * 100, 1)
            if expected
            else 0
        )

        if not window.is_active:
            status = "CLOSED"

        elif window.is_reopened:
            status = "REOPENED"

        elif timezone.now() >= window.deadline:
            status = "DEADLINE PASSED"

        elif (
            timezone.now()
            + timedelta(minutes=window.warning_minutes)
            >= window.deadline
        ):
            status = "CLOSING SOON"

        else:
            status = "OPEN"

        assessment_data.append({
            "window": window,
            "recorded": recorded,
            "expected": expected,
            "missing": max(expected - recorded, 0),
            "percentage": percentage,
            "status": status,
            "deadline_iso": window.deadline.isoformat(),
        })

    return render(
        request,
        "academic/marks_entry.html",
        {
            "students": students,
            "subjects": subjects,
            "assessment_types": assessment_types,
            "academic_years": academic_years,
            "terms": terms,
            "available_windows": open_windows,
            "available_years": available_years,
            "available_terms": available_terms,
            "available_assessment_types": available_assessment_types,

            # Exact Year + Term + Assessment combinations.
            "available_window_data": [
                {
                    "year": window.academic_year_id,
                    "term": window.term_id,
                    "assessment": window.assessment_type_id,
                }
                for window in open_windows
            ],
            "class_names": class_names,
            "streams": streams,
            "selected_student": selected_student,
            "selected_year": selected_year,
            "selected_term": selected_term,
            "selected_assessment_type": selected_assessment_type,
            "selected_subject": selected_subject,
            "selected_subjects": selected_subjects,
            "selected_class": selected_class,
            "selected_stream": selected_stream,
            "individual_assessments": individual_assessments,
            "class_students": class_students,
            "class_subject_marks": class_subject_marks,
            "all_subject_marks": all_subject_marks,
            "all_subjects_selected": all_subjects_selected,
            "can_enter_all_subjects": (
                is_admin_user(request.user)
                or (
                    selected_class
                    and is_class_teacher_for(
                        request.user,
                        selected_class,
                        selected_stream,
                    )
                )
            ),
            "mode": mode,
            "assessment_data": assessment_data,
            "can_configure": is_admin_user(request.user),
        }
    )


# ============================================================
# REPORT CARD CLASS LIST
# ============================================================


# ============================================================
# REPORT CARD CLASS LIST
# ============================================================

@login_required
@login_required

@login_required
@login_required
def report_cards(request):

    academic_years = AcademicYear.objects.all().order_by(
        "-year"
    )

    terms = Term.objects.all().order_by(
        "order"
    )

    assessment_types = AssessmentType.objects.filter(
        is_active=True
    ).order_by(
        "order",
        "name"
    )

    selected_class = request.GET.get(
        "class_name",
        ""
    ).strip()

    selected_stream = request.GET.get(
        "stream",
        ""
    ).strip()

    # ========================================================
    # STAFF-RESTRICTED CLASS AND STREAM OPTIONS
    # ========================================================

    class_names = get_staff_class_names(
        request.user
    )

    streams = get_staff_stream_names(
        request.user,
        selected_class
    )

    selected_year = None
    selected_term = None
    selected_assessment_type = None

    admission_no = request.GET.get(
        "admission_no",
        ""
    ).strip()

    student = None
    students = []

    branding = SchoolBranding.objects.first()

    # ========================================================
    # CHECK SELECTED CLASS / STREAM
    # ========================================================

    if selected_class:

        if not can_view_class(
            request.user,
            selected_class,
            selected_stream
        ):

            selected_class = ""
            selected_stream = ""

    # ========================================================
    # FIND LEARNER BY ADMISSION NUMBER
    # ========================================================

    if admission_no:

        student_queryset = (
            Student.objects
            .filter(
                admission_no__iexact=admission_no
            )
        )

        if not is_admin_user(request.user):

            student_queryset = (
                student_queryset.filter(
                    pk__in=get_staff_students(
                        request.user
                    ).values("pk")
                )
            )

        student = student_queryset.first()

        if student:

            selected_class = (
                student.class_name or ""
            )

            selected_stream = (
                get_student_stream_name(student)
            )

            if not can_view_class(
                request.user,
                selected_class,
                selected_stream
            ):

                student = None
                selected_class = ""
                selected_stream = ""

    # ========================================================
    # SELECTED YEAR
    # ========================================================

    year_id = request.GET.get(
        "year"
    )

    if year_id:

        selected_year = get_object_or_404(
            AcademicYear,
            id=year_id
        )

    # ========================================================
    # SELECTED TERM
    # ========================================================

    term_id = request.GET.get(
        "term"
    )

    if term_id:

        selected_term = get_object_or_404(
            Term,
            id=term_id
        )

    # ========================================================
    # SELECTED EXAMINATION
    # ========================================================

    assessment_type_id = request.GET.get(
        "assessment_type"
    )

    if assessment_type_id:

        selected_assessment_type = get_object_or_404(
            AssessmentType,
            id=assessment_type_id,
            is_active=True
        )

    # ========================================================
    # LOAD LEARNER
    # ========================================================

    if (
        student
        and selected_year
        and selected_term
        and selected_assessment_type
        and can_view_class(
            request.user,
            student.class_name,
            get_student_stream_name(student)
        )
    ):

        students = [student]

    # ========================================================
    # REPORT BOOK LOAD STATUS
    # ========================================================

    report_loaded = bool(
        student
        and selected_year
        and selected_term
        and selected_assessment_type
    )

    report_has_results = False

    if report_loaded:

        report_has_results = Assessment.objects.filter(
            student=student,
            academic_year=selected_year,
            term=selected_term,
            assessment_type=selected_assessment_type,
        ).exists()

    return render(
        request,
        "academic/report_cards.html",
        {
            "academic_years": academic_years,
            "terms": terms,
            "assessment_types": assessment_types,
            "class_names": class_names,
            "streams": streams,
            "selected_year": selected_year,
            "selected_term": selected_term,
            "selected_assessment_type": selected_assessment_type,
            "selected_class": selected_class,
            "selected_stream": selected_stream,
            "admission_no": admission_no,
            "student": student,
            "students": students,
            "branding": branding,
            "report_loaded": report_loaded,
            "report_has_results": report_has_results,
        }
    )


def class_list(request):

    academic_years = AcademicYear.objects.all().order_by(
        "-year"
    )

    terms = Term.objects.all().order_by(
        "order"
    )

    class_names = get_staff_class_names(
        request.user
    )

    selected_year = None
    selected_term = None
    selected_class = ""
    selected_stream = ""

    students = []

    branding = SchoolBranding.objects.first()

    # ========================================================
    # GET SELECTIONS
    # ========================================================

    year_id = request.GET.get(
        "year",
        ""
    ).strip()

    term_id = request.GET.get(
        "term",
        ""
    ).strip()

    selected_class = request.GET.get(
        "class_name",
        ""
    ).strip()

    selected_stream = request.GET.get(
        "stream",
        ""
    ).strip()

    # ========================================================
    # ACADEMIC YEAR
    # ========================================================

    if year_id:

        selected_year = get_object_or_404(
            AcademicYear,
            id=year_id
        )

    # ========================================================
    # TERM
    # ========================================================

    if term_id:

        selected_term = get_object_or_404(
            Term,
            id=term_id
        )

    # ========================================================
    # STREAMS AVAILABLE TO USER
    # ========================================================

    streams = get_staff_stream_names(
        request.user,
        selected_class
    )

    # ========================================================
    # LOAD CLASS
    # ========================================================

    if (
        selected_year
        and selected_term
        and selected_class
    ):

        if not can_view_class(
            request.user,
            selected_class,
            selected_stream
        ):

            selected_class = ""
            selected_stream = ""

        else:

            students_queryset = (
                get_staff_students(
                    request.user
                )
                .filter(
                    class_name=selected_class
                )
            )

            if selected_stream:

                students_queryset = (
                    students_queryset.filter(
                        stream=selected_stream
                    )
                )

            students = list(
                students_queryset.order_by(
                    "first_name",
                    "middle_name",
                    "last_name"
                )
            )

    # ========================================================
    # RENDER
    # ========================================================

    return render(
        request,
        "academic/class_list.html",
        {
            "academic_years": academic_years,
            "terms": terms,
            "class_names": class_names,
            "streams": streams,
            "selected_year": selected_year,
            "selected_term": selected_term,
            "selected_class": selected_class,
            "selected_stream": selected_stream,
            "students": students,
            "branding": branding,
        }
    )


@login_required
def student_report_card(
    request,
    student_id,
    academic_year_id,
    term_id
):

    student = get_object_or_404(
        Student,
        id=student_id
    )

    # ========================================================
    # ACCESS CONTROL
    # ========================================================

    if not is_admin_user(request.user):

        if not get_staff_students(
            request.user
        ).filter(
            pk=student.pk
        ).exists():

            return redirect(
                "academic:class_list"
            )

    academic_year = get_object_or_404(
        AcademicYear,
        id=academic_year_id
    )

    term = get_object_or_404(
        Term,
        id=term_id
    )

    assessment_type_id = request.GET.get(
        "assessment_type"
    )

    selected_assessment_type = None

    if assessment_type_id:

        selected_assessment_type = get_object_or_404(
            AssessmentType,
            id=assessment_type_id,
            is_active=True
        )

    subjects = Subject.objects.filter(
        is_active=True
    ).order_by("name")

    assessment_queryset = Assessment.objects.filter(
        student=student,
        academic_year=academic_year,
        term=term,
    ).select_related(
        "subject",
        "assessment_type",
    )

    if selected_assessment_type:

        assessment_queryset = (
            assessment_queryset.filter(
                assessment_type=selected_assessment_type
            )
        )

    assessments = list(
        assessment_queryset.order_by(
            "subject__name",
            "assessment_type__order",
        )
    )

    subject_results = []

    for subject in subjects:

        subject_assessments = [
            result
            for result in assessments
            if result.subject_id == subject.id
        ]

        if not subject_assessments:

            subject_results.append({
                "subject": subject,
                "assessments": [],
                "score": None,
                "grade": "___",
                "points": "___",
                "status": "blank",
            })

            continue

        if selected_assessment_type:

            result = subject_assessments[0]

            subject_results.append({
                "subject": subject,
                "assessments": subject_assessments,
                "score": float(result.score),
                "grade": result.grade(),
                "points": result.points(),
                "status": "recorded",
            })

        else:

            total = sum(
                float(result.score)
                for result in subject_assessments
            )

            average = (
                total /
                len(subject_assessments)
            )

            first_result = subject_assessments[0]

            subject_results.append({
                "subject": subject,
                "assessments": subject_assessments,
                "score": average,
                "grade": (
                    first_result.grade()
                    if average ==
                    float(first_result.score)
                    else "-"
                ),
                "points": "___",
                "status": "recorded",
            })

    recorded_scores = [
        item["score"]
        for item in subject_results
        if item["score"] is not None
    ]

    total = sum(
        recorded_scores
    )

    average = (
        total / len(recorded_scores)
        if recorded_scores
        else 0
    )

    class_students = get_staff_students(
        request.user
    ).filter(
        class_name=student.class_name
    )

    student_stream = get_student_stream_name(
        student
    )

    if student_stream:
        class_students = class_students.filter(
            stream=student_stream
        )

    learner_totals = []

    for learner in class_students:

        learner_assessments = Assessment.objects.filter(
            student=learner,
            academic_year=academic_year,
            term=term,
        )

        if selected_assessment_type:

            learner_assessments = (
                learner_assessments.filter(
                    assessment_type=selected_assessment_type
                )
            )

        learner_scores = list(
            learner_assessments.values_list(
                "score",
                flat=True
            )
        )

        if not learner_scores:
            continue

        learner_total = sum(
            float(score)
            for score in learner_scores
        )

        learner_average = (
            learner_total /
            len(learner_scores)
        )

        learner_totals.append({
            "student_id": learner.id,
            "average": learner_average,
            "total": learner_total,
        })

    learner_totals.sort(
        key=lambda item: item["average"],
        reverse=True
    )

    position = None

    for index, item in enumerate(
        learner_totals,
        start=1
    ):

        if item["student_id"] == student.id:

            position = index

            break

    class_size = len(
        learner_totals
    )

    context = {
        "student": student,
        "academic_year": academic_year,
        "term": term,
        "assessment_type": selected_assessment_type,
        "subjects": subjects,
        "results": assessments,
        "subject_results": subject_results,
        "total": total,
        "average": average,
        "position": position,
        "class_size": class_size,
        "student_stream": student_stream,
        "has_results": bool(recorded_scores),
    }

    return render(
        request,
        "academic/report_card.html",
        context
    )


@login_required
def academic_results(request):

    academic_years = AcademicYear.objects.all().order_by(
        "-year"
    )

    terms = Term.objects.all().order_by(
        "order"
    )

    assessment_types = AssessmentType.objects.filter(
        is_active=True
    ).order_by(
        "order",
        "name"
    )

    subjects = Subject.objects.filter(
        is_active=True
    ).order_by(
        "name"
    )

    selected_class = request.GET.get(
        "class_name",
        ""
    ).strip()

    selected_stream = request.GET.get(
        "stream",
        ""
    ).strip()

    class_names = get_staff_class_names(
        request.user
    )

    streams = get_staff_stream_names(
        request.user,
        selected_class
    )

    results = []

    selected_year = None
    selected_term = None
    selected_assessment_type = None
    selected_subject = None

    year_id = request.GET.get(
        "year"
    )

    term_id = request.GET.get(
        "term"
    )

    assessment_type_id = request.GET.get(
        "assessment_type"
    )

    subject_id = request.GET.get(
        "subject"
    )

    # ========================================================
    # CHECK SELECTED CLASS / STREAM
    # ========================================================

    class_is_allowed = True

    if selected_class:

        class_is_allowed = can_view_class(
            request.user,
            selected_class,
            selected_stream
        )

        if not class_is_allowed:

            selected_class = ""
            selected_stream = ""

    # ========================================================
    # SELECTED YEAR
    # ========================================================

    if year_id:

        selected_year = get_object_or_404(
            AcademicYear,
            id=year_id
        )

    # ========================================================
    # SELECTED TERM
    # ========================================================

    if term_id:

        selected_term = get_object_or_404(
            Term,
            id=term_id
        )

    # ========================================================
    # SELECTED ASSESSMENT TYPE
    # ========================================================

    if assessment_type_id:

        selected_assessment_type = get_object_or_404(
            AssessmentType,
            id=assessment_type_id,
            is_active=True
        )

    # ========================================================
    # SELECTED SUBJECT
    # ========================================================

    if subject_id and subject_id != "all":

        selected_subject = get_object_or_404(
            Subject,
            id=subject_id,
            is_active=True
        )

    # ========================================================
    # LOAD RESULTS
    # ========================================================

    if (
        selected_year
        and selected_term
        and selected_assessment_type
        and selected_class
        and class_is_allowed
    ):

        students = (
            get_staff_students(
                request.user
            )
            .filter(
                class_name=selected_class
            )
        )

        if selected_stream:

            students = students.filter(
                stream=selected_stream
            )

        for student in students:

            assessment_query = Assessment.objects.filter(
                student=student,
                academic_year=selected_year,
                term=selected_term,
                assessment_type=selected_assessment_type,
            ).select_related(
                "subject"
            )

            if selected_subject:

                assessment_query = (
                    assessment_query.filter(
                        subject=selected_subject
                    )
                )

            assessments = list(
                assessment_query.order_by(
                    "subject__name"
                )
            )

            total = sum(
                float(a.score)
                for a in assessments
            )

            average = (
                total / len(assessments)
                if assessments
                else 0
            )

            results.append({
                "student": student,
                "assessments": assessments,
                "total": total,
                "average": average,
            })

        results.sort(
            key=lambda item: item["average"],
            reverse=True
        )

        for position, result in enumerate(
            results,
            start=1
        ):

            result["position"] = position

    return render(
        request,
        "academic/academic_results.html",
        {
            "academic_years": academic_years,
            "terms": terms,
            "assessment_types": assessment_types,
            "subjects": subjects,
            "class_names": class_names,
            "streams": streams,
            "selected_year": selected_year,
            "selected_term": selected_term,
            "selected_assessment_type": selected_assessment_type,
            "selected_subject": selected_subject,
            "selected_class": selected_class,
            "selected_stream": selected_stream,
            "results": results,
        }
    )


@login_required
def performance_centre(request):
    """
    Navigation centre for learner, class, school performance,
    and report-card workflows.
    """
    return render(
        request,
        "academic/performance_centre.html",
        {
            "is_admin": is_admin_user(request.user),
            "branding": SchoolBranding.objects.first(),
        },
    )


@login_required
def performance_analysis(request):

    branding = (SchoolBranding.objects.filter(is_active=True).first() or SchoolBranding.objects.first())

    academic_years = AcademicYear.objects.all().order_by(
        "-year"
    )

    terms = Term.objects.all().order_by(
        "order"
    )

    assessment_types = AssessmentType.objects.filter(
        is_active=True
    ).order_by(
        "order",
        "name"
    )

    subjects = Subject.objects.filter(
        is_active=True
    ).order_by(
        "name"
    )

    selected_class = request.GET.get(
        "class_name",
        ""
    ).strip()

    selected_stream = request.GET.get(
        "stream",
        ""
    ).strip()

    class_names = get_staff_class_names(
        request.user
    )

    streams = get_staff_stream_names(
        request.user,
        selected_class
    )

    selected_year = None
    selected_term = None
    selected_assessment_type = None
    filter_error = ""

    year_id = request.GET.get(
        "year"
    )

    term_id = request.GET.get(
        "term"
    )

    assessment_type_id = request.GET.get(
        "assessment_type"
    )

    if selected_class:

        if not can_view_class(
            request.user,
            selected_class,
            selected_stream
        ):

            selected_class = ""
            selected_stream = ""

            streams = get_staff_stream_names(
                request.user,
                ""
            )

    # Apply class-specific subject permissions before calculating statistics.
    if selected_class:
        subjects = get_staff_subjects_for_class(
            request.user,
            selected_class,
            selected_stream,
        ).filter(
            is_active=True
        ).order_by("name")

    if year_id:
        if str(year_id).strip().isdigit():
            selected_year = AcademicYear.objects.filter(
                pk=year_id
            ).first()

        if selected_year:
            terms = Term.objects.filter(
                academic_year=selected_year
            ).order_by("order", "name")
        else:
            terms = Term.objects.none()
            filter_error = (
                "The selected academic year is invalid. "
                "Please choose a valid academic year."
            )

    if term_id:
        if str(term_id).strip().isdigit():
            selected_term = Term.objects.filter(
                pk=term_id
            ).first()

        if not selected_term and not filter_error:
            filter_error = (
                "The selected term is invalid. "
                "Please choose a valid term."
            )

    # Reject a term that does not belong to the selected year.
    if (
        selected_year
        and selected_term
        and selected_term.academic_year_id != selected_year.id
    ):
        filter_error = (
            "The selected term does not belong to the selected academic "
            "year. Please choose a term listed for that year."
        )
        selected_term = None

    if assessment_type_id:
        if str(assessment_type_id).strip().isdigit():
            selected_assessment_type = AssessmentType.objects.filter(
                pk=assessment_type_id,
                is_active=True,
            ).first()

        if not selected_assessment_type and not filter_error:
            filter_error = (
                "The selected assessment type is invalid or inactive. "
                "Please choose an available assessment type."
            )

    learner_count = 0
    class_average = 0
    highest_score = 0
    lowest_score = 0

    subject_analysis = []
    grade_distribution = []
    learner_analysis = []

    if (
        selected_year
        and selected_term
        and selected_term.academic_year_id == selected_year.id
        and selected_assessment_type
        and selected_class
        and can_view_class(
            request.user,
            selected_class,
            selected_stream
        )
    ):

        students = list(
            get_staff_students(
                request.user
            ).filter(
                class_name=selected_class
            )
        )

        if selected_stream:

            students = [
                student
                for student in students
                if get_student_stream_name(student)
                == selected_stream
            ]

        learner_count = len(
            students
        )

        assessments = Assessment.objects.filter(
            student__in=students,
            academic_year=selected_year,
            term=selected_term,
            assessment_type=selected_assessment_type,
            subject__in=subjects,
        ).select_related(
            "student",
            "subject"
        )

        overall_data = assessments.aggregate(
            average=Avg("score"),
            total=Sum("score"),
            count=Count("id"),
            highest=Max("score"),
            lowest=Min("score"),
        )

        class_average = (
            float(overall_data["average"])
            if overall_data["average"] is not None
            else 0
        )

        highest_score = (
            float(overall_data["highest"])
            if overall_data["highest"] is not None
            else 0
        )

        lowest_score = (
            float(overall_data["lowest"])
            if overall_data["lowest"] is not None
            else 0
        )

        for subject in subjects:

            subject_assessments = assessments.filter(
                subject=subject
            )

            data = subject_assessments.aggregate(
                average=Avg("score"),
                highest=Max("score"),
                lowest=Min("score"),
                count=Count("id"),
            )

            count = data["count"] or 0

            if count == 0:
                continue

            average = float(
                data["average"]
            )

            first_assessment = (
                subject_assessments.first()
            )

            subject_analysis.append({
                "subject": subject,
                "average": average,
                "highest": float(data["highest"]),
                "lowest": float(data["lowest"]),
                "count": count,
                "grade": (
                    first_assessment.grade()
                    if first_assessment
                    else "-"
                ),
            })

        grade_counts = {}

        for assessment in assessments:

            grade = assessment.grade()

            grade_counts[grade] = (
                grade_counts.get(
                    grade,
                    0
                ) + 1
            )

        for grade, count in grade_counts.items():

            grade_distribution.append({
                "grade": grade,
                "count": count,
            })

        for student in students:

            learner_scores = [
                float(a.score)
                for a in assessments
                if a.student_id == student.id
            ]

            learner_average = (
                sum(learner_scores)
                / len(learner_scores)
                if learner_scores
                else 0
            )

            learner_analysis.append({
                "student": student,
                "average": learner_average,
                "total": sum(learner_scores),
                "subjects": len(learner_scores),
            })

        learner_analysis.sort(
            key=lambda x: x["average"],
            reverse=True
        )

        for position, learner in enumerate(
            learner_analysis,
            start=1
        ):

            learner["position"] = position

    best_subject = None
    weakest_subject = None

    if subject_analysis:

        best_subject = max(
            subject_analysis,
            key=lambda item: item["average"]
        )

        weakest_subject = min(
            subject_analysis,
            key=lambda item: item["average"]
        )

    context = {

            "academic_years": academic_years,
            "terms": terms,
            "assessment_types": assessment_types,
            "subjects": subjects,
            "class_names": class_names,
            "streams": streams,
            "selected_year": selected_year,
            "selected_term": selected_term,
            "selected_assessment_type": selected_assessment_type,
            "filter_error": filter_error,
            "selected_class": selected_class,
            "selected_stream": selected_stream,
            "learner_count": learner_count,
            "class_average": class_average,
            "highest_score": highest_score,
            "lowest_score": lowest_score,
            "subject_analysis": subject_analysis,
            "grade_distribution": grade_distribution,
            "learner_analysis": learner_analysis,
            "best_subject": best_subject,
            "weakest_subject": weakest_subject,
            "branding": branding,
        }

    context.update(build_performance_analytics(
        request.user,
        selected_year,
        selected_term,
        selected_assessment_type,
        selected_class,
        selected_stream,
    ))

    return render(
        request,
        "academic/performance_analysis.html",
        context,
    )


















# ============================================================
# ASSESSMENT MANAGEMENT
# ============================================================

@login_required
def assessment_management(request):

    return_to_marks = request.POST.get("return_to_marks") == "1"

    academic_years = AcademicYear.objects.all().order_by("-year")

    terms = Term.objects.all().order_by("order")

    assessment_types = AssessmentType.objects.all().order_by(
        "order",
        "name"
    )

    windows = (
        AssessmentWindow.objects
        .select_related(
            "academic_year",
            "term",
            "assessment_type",
        )
        .filter(is_active=True)
        .order_by("-deadline")
    )

    # --------------------------------------------------------
    # CREATE / UPDATE ASSESSMENT WINDOW
    # --------------------------------------------------------

    if request.method == "POST":

        action = request.POST.get(
            "action",
            ""
        )

        # ----------------------------------------------------
        # SAVE WINDOW
        # ----------------------------------------------------

        if action == "save_window":

            if not is_admin_user(request.user):

                messages.error(
                    request,
                    "Only administrators can configure assessment deadlines."
                )

                return redirect(
                    "academic:marks_entry"
                    if return_to_marks
                    else "academic:assessment_management"
                )

            year_id = request.POST.get("academic_year")
            term_id = request.POST.get("term")
            assessment_type_id = request.POST.get(
                "assessment_type"
            )
            deadline_value = request.POST.get("deadline")
            warning_value = request.POST.get(
                "warning_minutes",
                "1440"
            )

            year = get_object_or_404(
                AcademicYear,
                id=year_id
            )

            term = get_object_or_404(
                Term,
                id=term_id
            )

            assessment_type = get_object_or_404(
                AssessmentType,
                id=assessment_type_id
            )

            if not deadline_value:

                messages.error(
                    request,
                    "Please provide a recording deadline."
                )

                return redirect(
                    "academic:marks_entry"
                    if return_to_marks
                    else "academic:assessment_management"
                )

            try:

                from datetime import datetime

                deadline = datetime.fromisoformat(
                    deadline_value
                )

                if timezone.is_naive(deadline):

                    deadline = timezone.make_aware(
                        deadline
                    )

                warning_minutes = int(
                    warning_value or 1440
                )

                AssessmentWindow.objects.update_or_create(

                    academic_year=year,

                    term=term,

                    assessment_type=assessment_type,

                    defaults={
                        "deadline": deadline,
                        "warning_minutes": max(
                            0,
                            warning_minutes
                        ),
                        "is_active": True,
                    }
                )

                messages.success(
                    request,
                    "Assessment recording window saved successfully."
                )

            except Exception as exc:

                messages.error(
                    request,
                    f"Could not save assessment window: {exc}"
                )

            return redirect(
                "academic:marks_entry"
                if return_to_marks
                else "academic:assessment_management"
            )

        # ----------------------------------------------------
        # REOPEN
        # ----------------------------------------------------

        if action == "reopen":

            if not is_admin_user(request.user):

                messages.error(
                    request,
                    "Only administrators can reopen assessments."
                )

                return redirect(
                    "academic:marks_entry"
                    if return_to_marks
                    else "academic:assessment_management"
                )

            window_id = request.POST.get(
                "window_id"
            )

            window = get_object_or_404(
                AssessmentWindow,
                id=window_id
            )

            window.is_reopened = True
            window.is_active = True
            window.save(
                update_fields=[
                    "is_reopened",
                    "is_active",
                    "updated_at",
                ]
            )

            messages.success(
                request,
                f"{window.assessment_type.name} has been reopened for recording."
            )

            return redirect(
                "academic:marks_entry"
                if return_to_marks
                else "academic:assessment_management"
            )

        # ----------------------------------------------------
        # CLOSE
        # ----------------------------------------------------

        if action == "close":

            if not is_admin_user(request.user):

                messages.error(
                    request,
                    "Only administrators can close assessments."
                )

                return redirect(
                    "academic:marks_entry"
                    if return_to_marks
                    else "academic:assessment_management"
                )

            window_id = request.POST.get(
                "window_id"
            )

            window = get_object_or_404(
                AssessmentWindow,
                id=window_id
            )

            window.is_reopened = False
            window.is_active = False

            window.save(
                update_fields=[
                    "is_reopened",
                    "is_active",
                    "updated_at",
                ]
            )

            messages.success(
                request,
                f"{window.assessment_type.name} has been closed."
            )

            return redirect(
                "academic:marks_entry"
                if return_to_marks
                else "academic:assessment_management"
            )

        # ----------------------------------------------------
        # UPLOAD MARKS
        # ----------------------------------------------------

        if action == "upload_marks":

            # ----------------------------------------------------
            # ONLY OPEN ASSESSMENTS MAY RECEIVE MARKS
            # ----------------------------------------------------

            window_id = request.POST.get(
                "upload_window"
            )

            upload_file = request.FILES.get(
                "marks_file"
            )

            if not window_id or not upload_file:

                messages.error(
                    request,
                    "Select an assessment and Excel file."
                )

                return redirect(
                    "academic:marks_entry"
                    if return_to_marks
                    else "academic:assessment_management"
                )

            window = get_object_or_404(
                AssessmentWindow,
                id=window_id
            )

            # ----------------------------------------------------
            # CHECK MARKS ENTRY PERMISSION
            # ----------------------------------------------------
            # can_enter_marks() requires the class name.
            # The Excel file contains learners from one selected
            # class, so determine the class from the first learner
            # row before processing the marks.

            try:
                from openpyxl import load_workbook

                permission_workbook = load_workbook(
                    upload_file,
                    read_only=True,
                    data_only=True
                )

                permission_sheet = permission_workbook.active

                first_learner_admission = None

                for permission_row in permission_sheet.iter_rows(
                    min_row=2,
                    values_only=True
                ):
                    if permission_row and permission_row[0]:
                        first_learner_admission = str(
                            permission_row[0]
                        ).strip()
                        break

                permission_workbook.close()

            except Exception:

                messages.error(
                    request,
                    "The Excel file could not be read. "
                    "Please use the generated marks template."
                )

                return redirect(
                    "academic:marks_entry"
                    if return_to_marks
                    else "academic:assessment_management"
                )

            if not first_learner_admission:

                messages.error(
                    request,
                    "The Excel file contains no learner records."
                )

                return redirect(
                    "academic:marks_entry"
                    if return_to_marks
                    else "academic:assessment_management"
                )

            permission_student = Student.objects.filter(
                admission_no__iexact=first_learner_admission
            ).first()

            if not permission_student:

                messages.error(
                    request,
                    f"Learner {first_learner_admission} was not found."
                )

                return redirect(
                    "academic:marks_entry"
                    if return_to_marks
                    else "academic:assessment_management"
                )

            if not can_enter_marks(
                request.user,
                permission_student.class_name
            ):

                messages.error(
                    request,
                    "You do not have permission to enter marks "
                    "for this class."
                )

                return redirect(
                    "academic:marks_entry"
                    if return_to_marks
                    else "academic:assessment_management"
                )

            # ----------------------------------------------------
            # ONLY OPEN ASSESSMENTS MAY RECEIVE MARKS
            # ----------------------------------------------------

            if not window.recording_open():

                messages.error(
                    request,
                    "Marks can only be recorded while this assessment "
                    "is OPEN. If the deadline has passed, the assessment "
                    "must be reopened first."
                )

                return redirect(
                    "academic:marks_entry"
                    if return_to_marks
                    else "academic:assessment_management"
                )

            upload_file = request.FILES.get(
                "marks_file"
            )

            if not window_id or not upload_file:

                messages.error(
                    request,
                    "Select an assessment and Excel file."
                )

                return redirect(
                    "academic:marks_entry"
                    if return_to_marks
                    else "academic:assessment_management"
                )

            window = get_object_or_404(
                AssessmentWindow,
                id=window_id
            )

            # Marks may ONLY be recorded while the assessment
            # window is OPEN. Reopened assessments are allowed
            # because recording_open() returns True when reopened.
            if not window.recording_open():

                messages.error(
                    request,
                    "Marks can only be recorded while this assessment is OPEN. "
                    "If the deadline has passed, the assessment must be reopened first."
                )

                return redirect(
                    "academic:marks_entry"
                    if return_to_marks
                    else "academic:assessment_management"
                )

            try:

                from openpyxl import load_workbook

                workbook = load_workbook(
                    upload_file,
                    data_only=True
                )

                sheet = workbook.active

                rows = list(
                    sheet.iter_rows(
                        values_only=True
                    )
                )

                if not rows:

                    raise ValueError(
                        "The Excel file is empty."
                    )

                headers = [
                    str(x).strip().lower()
                    if x is not None
                    else ""
                    for x in rows[0]
                ]

                # ------------------------------------------------
                # ACCEPT POPULATED MARKS TEMPLATE
                # ------------------------------------------------
                # The generated template contains:
                # Admission No, Learner Name, Subject, Mark
                #
                # Older templates may contain:
                # Admission No, Subject Code, Mark
                #
                # Support both formats.

                if "admission no" not in headers:
                    raise ValueError(
                        "Missing column: admission no"
                    )

                if "mark" not in headers:
                    raise ValueError(
                        "Missing column: mark"
                    )

                if (
                    "subject code" not in headers
                    and "subject" not in headers
                ):
                    raise ValueError(
                        "Missing subject column. "
                        "The Excel file must contain Subject "
                        "or Subject Code."
                    )

                admission_index = headers.index(
                    "admission no"
                )

                mark_index = headers.index(
                    "mark"
                )

                subject_code_index = (
                    headers.index("subject code")
                    if "subject code" in headers
                    else None
                )

                subject_name_index = (
                    headers.index("subject")
                    if "subject" in headers
                    else None
                )

                comment_index = (
                    headers.index(
                        "teacher comment"
                    )
                    if "teacher comment" in headers
                    else None
                )

                imported = 0
                errors = []

                for row_number, row in enumerate(
                    rows[1:],
                    start=2
                ):

                    admission = (
                        str(row[admission_index]).strip()
                        if row[admission_index] is not None
                        else ""
                    )

                    subject_code = ""

                    subject_name = ""

                    if (
                        subject_code_index is not None
                        and len(row) > subject_code_index
                        and row[subject_code_index] is not None
                    ):
                        subject_code = str(
                            row[subject_code_index]
                        ).strip()

                    if (
                        subject_name_index is not None
                        and len(row) > subject_name_index
                        and row[subject_name_index] is not None
                    ):
                        subject_name = str(
                            row[subject_name_index]
                        ).strip()

                    mark_value = row[mark_index]

                    if not admission and not subject_code:
                        continue

                    student = Student.objects.filter(
                        admission_no__iexact=admission
                    ).first()

                    if not student:

                        errors.append(
                            f"Row {row_number}: learner "
                            f"{admission} not found."
                        )

                        continue

                    # Find the subject by code when available.
                    # Otherwise find it by the populated subject name.

                    if subject_code:

                        subject = Subject.objects.filter(
                            code__iexact=subject_code,
                            is_active=True
                        ).first()

                    else:

                        subject = Subject.objects.filter(
                            name__iexact=subject_name,
                            is_active=True
                        ).first()

                    if not subject:

                        subject_reference = (
                            subject_code
                            or subject_name
                            or "blank subject"
                        )

                        errors.append(
                            f"Row {row_number}: subject "
                            f"{subject_reference} not found."
                        )

                        continue

                    try:

                        score = float(
                            mark_value
                        )

                    except (
                        TypeError,
                        ValueError
                    ):

                        errors.append(
                            f"Row {row_number}: invalid mark."
                        )

                        continue

                    if score < 0 or score > 100:

                        errors.append(
                            f"Row {row_number}: mark "
                            f"must be between 0 and 100."
                        )

                        continue

                    comment = ""

                    if (
                        comment_index is not None
                        and len(row) > comment_index
                        and row[comment_index] is not None
                    ):

                        comment = str(
                            row[comment_index]
                        ).strip()

                    Assessment.objects.update_or_create(

                        student=student,

                        academic_year=window.academic_year,

                        term=window.term,

                        subject=subject,

                        assessment_type=window.assessment_type,

                        defaults={
                            "score": score,
                            "teacher_comment": comment,
                        }
                    )

                    imported += 1

                if errors:

                    messages.warning(
                        request,
                        f"{imported} marks imported. "
                        f"{len(errors)} rows rejected: "
                        + " | ".join(errors[:8])
                    )

                else:

                    messages.success(
                        request,
                        f"{imported} marks imported successfully."
                    )

            except ImportError:

                messages.error(
                    request,
                    "Excel upload requires openpyxl. "
                    "Run: python -m pip install openpyxl"
                )

            except Exception as exc:

                messages.error(
                    request,
                    f"Upload failed: {exc}"
                )

            return redirect(
                "academic:marks_entry"
                if return_to_marks
                else "academic:assessment_management"
            )

    # --------------------------------------------------------
    # BUILD PROGRESS
    # --------------------------------------------------------

    assessment_data = []

    for window in windows:

        total_students = Student.objects.filter(
            class_name__isnull=False
        ).exclude(
            class_name=""
        ).count()

        recorded = Assessment.objects.filter(
            academic_year=window.academic_year,
            term=window.term,
            assessment_type=window.assessment_type,
        ).count()

        subjects_count = Subject.objects.filter(
            is_active=True
        ).count()

        expected = (
            total_students *
            subjects_count
        )

        percentage = (
            round(
                (recorded / expected) * 100,
                1
            )
            if expected
            else 0
        )

        if not window.is_active:

            status = "CLOSED"

        elif window.is_reopened:

            status = "REOPENED"

        elif timezone.now() >= window.deadline:

            status = "DEADLINE PASSED"

        elif (
            timezone.now() +
            timedelta(
                minutes=window.warning_minutes
            )
            >= window.deadline
        ):

            status = "CLOSING SOON"

        else:

            status = "OPEN"

        assessment_data.append({
            "window": window,
            "recorded": recorded,
            "expected": expected,
            "missing": max(
                expected - recorded,
                0
            ),
            "percentage": percentage,
            "status": status,
            "deadline_iso": window.deadline.isoformat(),
        })

    return render(
        request,
        "academic/assessment_management.html",
        {
            "academic_years": academic_years,
            "terms": terms,
            "assessment_types": assessment_types,
            "assessment_data": assessment_data,
            "can_configure": is_admin_user(
                request.user
            ),
        }
    )


# ============================================================
# DOWNLOAD MARKS TEMPLATE
# ============================================================

@login_required
def assessment_marks_template(request):

    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, Protection
        from openpyxl.worksheet.datavalidation import DataValidation

        windows = AssessmentWindow.objects.filter(
            is_active=True
        ).select_related(
            "academic_year",
            "term",
            "assessment_type",
        ).order_by(
            "-deadline"
        )

        selected_window_id = request.GET.get(
            "assessment"
        )

        selected_class = request.GET.get(
            "class_name"
        )

        # --------------------------------------------------------
        # SHOW SELECTION PAGE
        # --------------------------------------------------------

        if not selected_window_id or not selected_class:

            classes = get_staff_class_names(
                request.user
            )

            return render(
                request,
                "academic/assessment_marks_template.html",
                {
                    "windows": windows,
                    "classes": classes,
                    "selected_window_id": selected_window_id,
                    "selected_class": selected_class,
                    "branding": SchoolBranding.objects.filter(
                        is_active=True
                    ).first(),
                }
            )

        # --------------------------------------------------------
        # LOAD SELECTED ASSESSMENT
        # --------------------------------------------------------

        window = get_object_or_404(
            AssessmentWindow.objects.select_related(
                "academic_year",
                "term",
                "assessment_type",
            ),
            id=selected_window_id,
        )

        # Only open assessments can generate recording templates.
        if not window.recording_open():

            messages.error(
                request,
                "An Excel marks template can only be generated "
                "while the assessment is OPEN."
            )

            return redirect(
                "academic:assessment_marks_template"
            )

        # --------------------------------------------------------
        # CHECK CLASS ACCESS
        # --------------------------------------------------------

        allowed_classes = list(
            get_staff_class_names(request.user)
        )

        if selected_class not in allowed_classes:

            messages.error(
                request,
                "You do not have permission to access this class."
            )

            return redirect(
                "academic:assessment_marks_template"
            )

        # --------------------------------------------------------
        # GET LEARNERS
        # --------------------------------------------------------

        students = Student.objects.filter(
            class_name=selected_class
        ).order_by(
            "admission_no",
            "first_name",
            "middle_name",
            "last_name",
        )

        # --------------------------------------------------------
        # GET ACTIVE SUBJECTS
        # --------------------------------------------------------

        subjects = Subject.objects.filter(
            is_active=True
        ).order_by(
            "name"
        )

        if not students.exists():

            messages.error(
                request,
                f"No learners were found in {selected_class}."
            )

            return redirect(
                "academic:assessment_marks_template"
            )

        if not subjects.exists():

            messages.error(
                request,
                "No active subjects are configured."
            )

            return redirect(
                "academic:assessment_marks_template"
            )

        # --------------------------------------------------------
        # CREATE WORKBOOK
        # --------------------------------------------------------

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Marks Upload"

        headers = [
            "Admission No",
            "Learner Name",
            "Subject",
            "Mark",
        ]

        sheet.append(headers)

        # Header formatting
        for cell in sheet[1]:

            cell.font = Font(
                bold=True
            )

            cell.protection = Protection(
                locked=True
            )

        # --------------------------------------------------------
        # POPULATE LEARNERS + SUBJECTS
        # --------------------------------------------------------

        row_number = 2

        for student in students:

            learner_name = " ".join(
                part
                for part in [
                    student.first_name,
                    student.middle_name,
                    student.last_name,
                ]
                if part
            ).strip()

            for subject in subjects:

                sheet.cell(
                    row=row_number,
                    column=1,
                    value=student.admission_no
                )

                sheet.cell(
                    row=row_number,
                    column=2,
                    value=learner_name
                )

                sheet.cell(
                    row=row_number,
                    column=3,
                    value=subject.name
                )

                mark_cell = sheet.cell(
                    row=row_number,
                    column=4,
                    value=""
                )

                # Only Mark is editable.
                for column in range(1, 4):

                    sheet.cell(
                        row=row_number,
                        column=column
                    ).protection = Protection(
                        locked=True
                    )

                mark_cell.protection = Protection(
                    locked=False
                )

                row_number += 1

        # --------------------------------------------------------
        # MARK VALIDATION: 0 - 100
        # --------------------------------------------------------

        validation = DataValidation(
            type="decimal",
            operator="between",
            formula1="0",
            formula2="100",
            allow_blank=True,
        )

        validation.error = (
            "Enter a mark between 0 and 100."
        )

        validation.errorTitle = "Invalid Mark"

        validation.prompt = (
            "Enter the learner's mark from 0 to 100."
        )

        validation.promptTitle = "Mark"

        sheet.add_data_validation(
            validation
        )

        if row_number > 2:

            validation.add(
                f"D2:D{row_number - 1}"
            )

        # --------------------------------------------------------
        # LOCK LEARNER INFORMATION
        # --------------------------------------------------------

        sheet.protection.sheet = True
        sheet.protection.password = "LUHAN_MARKS"

        sheet.freeze_panes = "D2"

        sheet.auto_filter.ref = (
            f"A1:D{row_number - 1}"
        )

        # Column widths
        sheet.column_dimensions["A"].width = 18
        sheet.column_dimensions["B"].width = 32
        sheet.column_dimensions["C"].width = 28
        sheet.column_dimensions["D"].width = 15

        # --------------------------------------------------------
        # INSTRUCTIONS SHEET
        # --------------------------------------------------------

        instructions = workbook.create_sheet(
            "Instructions"
        )

        instructions.append([
            "ASSESSMENT MARKS ENTRY"
        ])

        instructions.append([
            "School",
            (
                SchoolBranding.objects.filter(
                    is_active=True
                ).values_list(
                    "school_name",
                    flat=True
                ).first()
                or "School ERP"
            )
        ])

        instructions.append([
            "Assessment",
            window.assessment_type.name
        ])

        instructions.append([
            "Academic Year",
            str(window.academic_year)
        ])

        instructions.append([
            "Term",
            str(window.term)
        ])

        instructions.append([
            "Grade/Class",
            selected_class
        ])

        instructions.append([])

        instructions.append([
            "IMPORTANT"
        ])

        instructions.append([
            "Only enter marks in the Mark column."
        ])

        instructions.append([
            "Marks must be between 0 and 100."
        ])

        instructions.append([
            "Do not change Admission No, Learner Name or Subject."
        ])

        instructions.column_dimensions[
            "A"
        ].width = 25

        instructions.column_dimensions[
            "B"
        ].width = 70

        # --------------------------------------------------------
        # RESPONSE
        # --------------------------------------------------------

        response = HttpResponse(
            content_type=(
                "application/vnd.openxmlformats-officedocument."
                "spreadsheetml.sheet"
            )
        )

        safe_class = (
            str(selected_class)
            .replace(" ", "_")
            .replace("/", "-")
        )

        safe_assessment = (
            str(window.assessment_type.name)
            .replace(" ", "_")
            .replace("/", "-")
        )

        response[
            "Content-Disposition"
        ] = (
            'attachment; '
            f'filename="{safe_class}_{safe_assessment}_Marks.xlsx"'
        )

        workbook.save(
            response
        )

        return response

    except ImportError:

        return HttpResponse(
            "Excel generation requires openpyxl. "
            "Run: python -m pip install openpyxl",
            status=500
        )




@login_required
def performance_comparison(request):
    """
    Explicitly compare two selected assessment periods.

    Modes:
      school  - administrators only
      class   - class/stream permitted by existing staff access rules
      learner - learner visible through get_staff_students()

    No year, term, or assessment type is automatically selected.
    """
    from django.core.exceptions import PermissionDenied
    from django.db.models import Avg, Count
    from django.shortcuts import render
    from fees.models import AcademicYear, Term
    from .models import Assessment, AssessmentType

    admin_user = is_admin_user(request.user)

    years = AcademicYear.objects.all().order_by("-year")
    terms = Term.objects.select_related("academic_year").all()
    types = AssessmentType.objects.filter(
        is_active=True
    ).order_by("order", "name")

    mode = request.GET.get("mode", "").strip()
    if not mode:
        mode = "school" if admin_user else "class"

    if mode not in {"school", "class", "learner"}:
        mode = "class"

    if mode == "school" and not admin_user:
        raise PermissionDenied(
            "Whole-school performance comparison is restricted to administrators."
        )

    class_name = request.GET.get("class_name", "").strip()
    stream_name = request.GET.get("stream", "").strip()
    admission_no = request.GET.get("admission_no", "").strip()

    # The same scope is applied to both sides of the comparison.
    scoped_students = None
    scope_error = ""

    if mode == "class":
        if class_name and stream_name:
            if not can_view_class(request.user, class_name, stream_name):
                raise PermissionDenied(
                    "You do not have access to this class and stream."
                )

            scoped_students = get_staff_students(request.user).filter(
                class_name=class_name
            )
            stream_student_ids = [
                student.pk
                for student in scoped_students
                if str(get_student_stream_name(student) or "").strip().casefold()
                == stream_name.casefold()
            ]
            scoped_students = scoped_students.filter(
                pk__in=stream_student_ids
            )
        else:
            scope_error = "Enter the class and stream you want to compare."

    elif mode == "learner":
        if admission_no:
            scoped_students = get_staff_students(request.user).filter(
                admission_no__iexact=admission_no
            )
            if not scoped_students.exists():
                scope_error = (
                    "The learner was not found among the learners you are "
                    "permitted to access."
                )
        else:
            scope_error = "Enter the learner's admission number."

    # No default periods or types: every value comes from the user.
    fields = {
        "current_year": request.GET.get("current_year", "").strip(),
        "current_term": request.GET.get("current_term", "").strip(),
        "current_type": request.GET.get("current_type", "").strip(),
        "compare_year": request.GET.get("compare_year", "").strip(),
        "compare_term": request.GET.get("compare_term", "").strip(),
        "compare_type": request.GET.get("compare_type", "").strip(),
    }

    submitted = request.GET.get("compare") == "1"
    errors = []
    result = None
    subject_rows = []
    current_label = ""
    compare_label = ""

    def selected_object(model, raw_id, label, queryset=None):
        if not raw_id:
            errors.append(f"Select {label}.")
            return None
        qs = queryset if queryset is not None else model.objects.all()
        try:
            obj = qs.filter(pk=int(raw_id)).first()
        except (TypeError, ValueError):
            obj = None
        if obj is None:
            errors.append(f"The selected {label} is invalid.")
        return obj

    if submitted:
        current_year = selected_object(
            AcademicYear, fields["current_year"], "current academic year", years
        )
        compare_year = selected_object(
            AcademicYear, fields["compare_year"], "comparison academic year", years
        )
        current_term = selected_object(
            Term, fields["current_term"], "current term", terms
        )
        compare_term = selected_object(
            Term, fields["compare_term"], "comparison term", terms
        )
        current_type = selected_object(
            AssessmentType, fields["current_type"], "current assessment type", types
        )
        compare_type = selected_object(
            AssessmentType, fields["compare_type"], "comparison assessment type", types
        )

        if current_year and current_term:
            if current_term.academic_year_id != current_year.pk:
                errors.append("The current term must belong to the current academic year.")

        if compare_year and compare_term:
            if compare_term.academic_year_id != compare_year.pk:
                errors.append("The comparison term must belong to the comparison academic year.")

        if scope_error:
            errors.append(scope_error)

        if mode == "class" and scoped_students is not None and not scoped_students.exists():
            errors.append("No learners were found for the selected class and stream.")

        if (
            not errors
            and current_year and compare_year
            and current_term and compare_term
            and current_type and compare_type
            and scoped_students is not None if mode != "school" else True
        ):
            base = Assessment.objects.all()

            if mode == "class" or mode == "learner":
                base = base.filter(student__in=scoped_students)

            current_qs = base.filter(
                academic_year=current_year,
                term=current_term,
                assessment_type=current_type,
            )
            compare_qs = base.filter(
                academic_year=compare_year,
                term=compare_term,
                assessment_type=compare_type,
            )

            current_stats = current_qs.aggregate(
                average=Avg("score"),
                marks_count=Count("pk"),
                learners_count=Count("student_id", distinct=True),
            )
            compare_stats = compare_qs.aggregate(
                average=Avg("score"),
                marks_count=Count("pk"),
                learners_count=Count("student_id", distinct=True),
            )

            current_average = current_stats["average"]
            compare_average = compare_stats["average"]

            difference = None
            if current_average is not None and compare_average is not None:
                difference = round(
                    float(current_average) - float(compare_average), 2
                )

            current_subjects = {
                row["subject_id"]: row
                for row in current_qs.values(
                    "subject_id", "subject__name"
                ).annotate(
                    average=Avg("score"),
                    marks_count=Count("pk"),
                )
            }
            compare_subjects = {
                row["subject_id"]: row
                for row in compare_qs.values(
                    "subject_id", "subject__name"
                ).annotate(
                    average=Avg("score"),
                    marks_count=Count("pk"),
                )
            }

            all_subject_ids = set(current_subjects) | set(compare_subjects)
            for subject_id in sorted(
                all_subject_ids,
                key=lambda sid: (
                    (current_subjects.get(sid) or compare_subjects.get(sid))
                    ["subject__name"] or ""
                ).casefold(),
            ):
                cur = current_subjects.get(subject_id)
                old = compare_subjects.get(subject_id)
                cur_avg = cur["average"] if cur else None
                old_avg = old["average"] if old else None
                subject_difference = None
                if cur_avg is not None and old_avg is not None:
                    subject_difference = round(
                        float(cur_avg) - float(old_avg), 2
                    )

                subject_rows.append({
                    "name": (
                        cur["subject__name"] if cur
                        else old["subject__name"]
                    ),
                    "current_average": cur_avg,
                    "compare_average": old_avg,
                    "difference": subject_difference,
                    "current_count": cur["marks_count"] if cur else 0,
                    "compare_count": old["marks_count"] if old else 0,
                })

            current_label = (
                f"{current_year} ? {current_term} ? {current_type}"
            )
            compare_label = (
                f"{compare_year} ? {compare_term} ? {compare_type}"
            )

            result = {
                "current_average": current_average,
                "compare_average": compare_average,
                "difference": difference,
                "current_marks_count": current_stats["marks_count"],
                "compare_marks_count": compare_stats["marks_count"],
                "current_learners_count": current_stats["learners_count"],
                "compare_learners_count": compare_stats["learners_count"],
            }

    context = {
        "branding": SchoolBranding.objects.first(),
        "is_admin": admin_user,
        "mode": mode,
        "years": years,
        "terms": terms,
        "assessment_types": types,
        "fields": fields,
        "class_name": class_name,
        "stream_name": stream_name,
        "admission_no": admission_no,
        "submitted": submitted,
        "errors": errors,
        "result": result,
        "subject_rows": subject_rows,
        "current_label": current_label,
        "compare_label": compare_label,
        "scope_error": scope_error,
    }
    return render(
        request,
        "academic/performance_comparison.html",
        context,
    )


@login_required
def whole_school_performance_analysis(request):
    """Administrator-only school-wide academic performance analysis."""
    from django.core.exceptions import PermissionDenied
    from django.db.models import Avg, Count, Max, Min
    from django.shortcuts import get_object_or_404
    from fees.models import AcademicYear, Term
    from .models import Assessment, AssessmentType

    if not is_admin_user(request.user):
        raise PermissionDenied(
            "Whole-school performance analysis is restricted to administrators."
        )

    branding = (SchoolBranding.objects.filter(is_active=True).first() or SchoolBranding.objects.first())
    academic_years = AcademicYear.objects.all().order_by("-year")
    selected_year_id = request.GET.get("year", "").strip()
    selected_term_id = request.GET.get("term", "").strip()
    selected_type_id = request.GET.get("assessment_type", "").strip()

    selected_year = None
    selected_term = None
    selected_assessment_type = None
    filter_error = ""

    if selected_year_id:
        if selected_year_id.isdigit():
            selected_year = AcademicYear.objects.filter(
                pk=selected_year_id
            ).first()

        if not selected_year:
            filter_error = (
                "The selected academic year is invalid. "
                "Please choose a valid academic year."
            )
    else:
        selected_year = academic_years.first()

    terms = Term.objects.none()
    if selected_year:
        terms = Term.objects.filter(
            academic_year=selected_year
        ).order_by("order", "name")

        if selected_term_id:
            if selected_term_id.isdigit():
                selected_term = terms.filter(
                    pk=selected_term_id
                ).first()

            if not selected_term and not filter_error:
                filter_error = (
                    "The selected term is invalid or does not belong "
                    "to the selected academic year."
                )
        else:
            selected_term = terms.order_by("-order", "-pk").first()
    elif selected_term_id and not filter_error:
        filter_error = (
            "Choose a valid academic year before selecting a term."
        )

    assessment_types = AssessmentType.objects.filter(
        is_active=True
    ).order_by("order", "name")

    if selected_type_id:
        if selected_type_id.isdigit():
            selected_assessment_type = AssessmentType.objects.filter(
                pk=selected_type_id,
                is_active=True,
            ).first()

        if not selected_assessment_type and not filter_error:
            filter_error = (
                "The selected assessment type is invalid or inactive. "
                "Please choose an available assessment type."
            )

    empty_summary = {
        "average": None,
        "marks_count": 0,
        "learners_count": 0,
        "highest": None,
        "lowest": None,
    }
    summary = empty_summary.copy()
    assessment_comparison = []
    term_comparison = []
    subject_rankings = []
    class_rankings = []

    if selected_year:
        # Require the stored academic year and the term's academic year
        # to agree so inconsistent legacy records are excluded.
        year_marks = Assessment.objects.filter(
            academic_year=selected_year,
            term__academic_year=selected_year,
        ).select_related(
            "student", "subject", "term", "assessment_type"
        )

        current_marks = year_marks
        if selected_term:
            current_marks = current_marks.filter(term=selected_term)
        if selected_assessment_type:
            current_marks = current_marks.filter(
                assessment_type=selected_assessment_type
            )

        totals = current_marks.aggregate(
            average=Avg("score"),
            marks_count=Count("pk"),
            highest=Max("score"),
            lowest=Min("score"),
        )
        summary.update(totals)
        summary["learners_count"] = current_marks.order_by().values(
            "student_id"
        ).distinct().count()

        # Assessment types are compared inside the selected term.
        # Do not apply the selected type filter to this comparison.
        type_marks = year_marks
        if selected_term:
            type_marks = type_marks.filter(term=selected_term)

        assessment_comparison = list(
            type_marks.values(
                "assessment_type__name"
            ).annotate(
                average=Avg("score"),
                marks_count=Count("pk"),
                learners_count=Count("student_id", distinct=True),
            ).order_by("-average", "assessment_type__name")
        )

        # Term comparison spans the year. The chosen term must not
        # narrow this queryset; the assessment-type filter still applies.
        term_marks = year_marks
        if selected_assessment_type:
            term_marks = term_marks.filter(
                assessment_type=selected_assessment_type
            )

        term_comparison = list(
            term_marks.values(
                "term__name", "term__order"
            ).annotate(
                average=Avg("score"),
                marks_count=Count("pk"),
                learners_count=Count("student_id", distinct=True),
            ).order_by("term__order", "term__name")
        )

        subject_rankings = list(
            current_marks.values(
                "subject__name"
            ).annotate(
                average=Avg("score"),
                marks_count=Count("pk"),
                learners_count=Count("student_id", distinct=True),
            ).order_by("-average", "subject__name")
        )

        class_rankings = list(
            current_marks.values(
                "student__class_name"
            ).annotate(
                average=Avg("score"),
                marks_count=Count("pk"),
                learners_count=Count("student_id", distinct=True),
            ).order_by("-average", "student__class_name")
        )

    context = {
        "branding": branding,
        "academic_years": academic_years,
        "terms": terms,
        "assessment_types": assessment_types,
        "selected_year": selected_year,
        "selected_term": selected_term,
        "selected_assessment_type": selected_assessment_type,
        "summary": summary,
        "assessment_comparison": assessment_comparison,
        "term_comparison": term_comparison,
        "subject_rankings": subject_rankings,
        "class_rankings": class_rankings,
        "filter_error": filter_error,
    }
    return render(
        request,
        "academic/whole_school_performance.html",
        context,
    )



@login_required
def learner_performance_analysis(request):
    """Compare a permitted learner's recorded results across assessments and terms."""

    branding = (SchoolBranding.objects.filter(is_active=True).first() or SchoolBranding.objects.first())
    academic_years = AcademicYear.objects.all().order_by("-year")
    assessment_types = AssessmentType.objects.filter(
        is_active=True
    ).order_by("order", "name")

    admission_no = request.GET.get("admission_no", "").strip()
    year_id = request.GET.get("year", "").strip()
    term_id = request.GET.get("term", "").strip()
    assessment_type_id = request.GET.get("assessment_type", "").strip()

    student = None
    selected_year = None
    selected_term = None
    selected_assessment_type = None
    terms = Term.objects.none()
    assessment_comparison = []
    term_comparison = []
    subject_comparison = []
    overall_average = None
    recorded_subject_count = 0
    error_message = ""

    if year_id:
        if year_id.isdigit():
            selected_year = AcademicYear.objects.filter(
                pk=year_id
            ).first()

        if selected_year:
            terms = Term.objects.filter(
                academic_year=selected_year
            ).order_by("order", "name")
        else:
            terms = Term.objects.none()
            error_message = (
                "The selected academic year is invalid. "
                "Please choose a valid academic year."
            )

    if term_id:
        if selected_year:
            term_query = Term.objects.filter(
                academic_year=selected_year
            )
        elif not year_id:
            term_query = Term.objects.all()
        else:
            term_query = Term.objects.none()

        if term_id.isdigit():
            selected_term = term_query.filter(pk=term_id).first()

        if not selected_term and not error_message:
            error_message = (
                "The selected term is invalid or does not belong "
                "to the selected academic year."
            )

    if assessment_type_id:
        if assessment_type_id.isdigit():
            selected_assessment_type = AssessmentType.objects.filter(
                pk=assessment_type_id,
                is_active=True,
            ).first()

        if not selected_assessment_type and not error_message:
            error_message = (
                "The selected assessment type is invalid or inactive. "
                "Please choose an available assessment type."
            )

    if admission_no:
        # This queryset is the staff-access boundary for learner records.
        student = get_staff_students(request.user).filter(
            admission_no__iexact=admission_no
        ).first()

        if student is None:
            error_message = (
                "No learner with that admission number was found "
                "in the records you are permitted to access."
            )

    if student and selected_year:
        marks = Assessment.objects.filter(
            student=student,
            academic_year=selected_year,
        ).select_related("term", "subject", "assessment_type")

        # Compare assessment types within the selected term.
        if selected_term:
            current_marks = marks.filter(term=selected_term)

            for assessment_type in assessment_types:
                rows = current_marks.filter(assessment_type=assessment_type)
                stats = rows.aggregate(
                    average=Avg("score"),
                    count=Count("id"),
                    subjects=Count("subject_id", distinct=True),
                )
                assessment_comparison.append({
                    "name": assessment_type.name,
                    "average": (
                        round(float(stats["average"]), 2)
                        if stats["average"] is not None else None
                    ),
                    "count": stats["count"],
                    "subjects": stats["subjects"],
                })

            subject_marks = current_marks
            if selected_assessment_type:
                subject_marks = subject_marks.filter(
                    assessment_type=selected_assessment_type
                )

            subject_stats = list(
                subject_marks.values("subject__name").annotate(
                    average=Avg("score"),
                    count=Count("id"),
                ).order_by("subject__name")
            )
            subject_comparison = [{
                "name": item["subject__name"],
                "average": (
                    round(float(item["average"]), 2)
                    if item["average"] is not None else None
                ),
                "count": item["count"],
            } for item in subject_stats]

            summary = subject_marks.aggregate(
                average=Avg("score"),
                subjects=Count("subject_id", distinct=True),
            )
            overall_average = (
                round(float(summary["average"]), 2)
                if summary["average"] is not None else None
            )
            recorded_subject_count = summary["subjects"] or 0

        # Compare terms within the selected academic year.
        term_marks = marks
        if selected_assessment_type:
            term_marks = term_marks.filter(
                assessment_type=selected_assessment_type
            )

        term_stats = list(
            term_marks.values(
                "term_id", "term__name", "term__order"
            ).annotate(
                average=Avg("score"),
                count=Count("id"),
                subjects=Count("subject_id", distinct=True),
            ).order_by("term__order", "term_id")
        )

        term_comparison = [{
            "name": item["term__name"],
            "order": item["term__order"],
            "average": (
                round(float(item["average"]), 2)
                if item["average"] is not None else None
            ),
            "count": item["count"],
            "subjects": item["subjects"],
        } for item in term_stats]

    return render(
        request,
        "academic/learner_performance_analysis.html",
        {
            "branding": branding,
            "academic_years": academic_years,
            "terms": terms,
            "assessment_types": assessment_types,
            "admission_no": admission_no,
            "student": student,
            "selected_year": selected_year,
            "selected_term": selected_term,
            "selected_assessment_type": selected_assessment_type,
            "assessment_comparison": assessment_comparison,
            "term_comparison": term_comparison,
            "subject_comparison": subject_comparison,
            "overall_average": overall_average,
            "recorded_subject_count": recorded_subject_count,
            "error_message": error_message,
        },
    )

