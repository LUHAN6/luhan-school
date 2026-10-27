from datetime import timedelta
from django.test import TestCase
from django.core.exceptions import ValidationError
from django.utils import timezone
from .models import Vehicle, TransportStaff, TransportRoute, RouteStop, LearnerTransport, TransportTrip, VehicleInspection


class TransportModelTests(TestCase):
    def setUp(self):
        self.vehicle = Vehicle.objects.create(registration_number="KXX 123A", seating_capacity=14)
        self.driver = TransportStaff.objects.create(full_name="Test Driver", role="DRIVER")
        self.route = TransportRoute.objects.create(name="Test Route", code="TR1")
        self.stop = RouteStop.objects.create(route=self.route, name="Test Stop", sequence=1)
    def test_stop_must_belong_to_selected_route(self):
        other = TransportRoute.objects.create(name="Other Route", code="TR2")
        assignment = LearnerTransport(
            learner_admission_no="TEST001", learner_name="Test Learner",
            route=other, stop=self.stop
        )
        with self.assertRaises(ValidationError):
            assignment.full_clean()
    def test_expired_vehicle_documents_flag(self):
        self.vehicle.insurance_expiry = timezone.localdate() - timedelta(days=1)
        self.vehicle.save()
        self.assertTrue(self.vehicle.documents_expired)
    def test_inspection_requires_all_safety_items(self):
        trip = TransportTrip.objects.create(route=self.route, direction="AM", vehicle=self.vehicle, driver=self.driver)
        inspection = VehicleInspection.objects.create(trip=trip, result="PASS", brakes_ok=True, tyres_ok=True)
        self.assertFalse(inspection.dispatch_allowed)


from decimal import Decimal
from django.contrib.auth import get_user_model
from django.contrib.auth.models import Permission
from django.urls import reverse
from accounts.models import Role, UserProfile
from .models import (
    LearnerTransport,
    TransportCharge,
    TransportPayment,
    TransportAuditEvent,
)


class TransportPaymentVerificationTests(TestCase):
    def setUp(self):
        self.user_model = get_user_model()
        self.staff = self.user_model.objects.create_user(
            username="transport_test_staff",
            password="test-password-not-for-production",
        )
        self.role = Role.objects.create(
            name="Transport Test Role",
            is_active=True,
        )
        UserProfile.objects.update_or_create(
            user=self.staff,
            defaults={"role": "STAFF", "custom_role": self.role},
        )

        route = TransportRoute.objects.create(
            name="Payment Test Route", code="PAYTEST"
        )
        stop = RouteStop.objects.create(
            route=route, name="Payment Test Stop", sequence=1
        )
        learner = LearnerTransport.objects.create(
            learner_admission_no="PAYTEST001",
            learner_name="Payment Test Learner",
            route=route,
            stop=stop,
        )
        self.charge = TransportCharge.objects.create(
            learner=learner,
            period_label="Test Term",
            amount=Decimal("1000.00"),
            status="DUE",
        )
        self.payment = TransportPayment.objects.create(
            charge=self.charge,
            amount=Decimal("600.00"),
            transaction_reference="PAYTEST-REF-001",
        )
        self.url = reverse(
            "transport_pro:verify_payment",
            kwargs={"pk": self.payment.pk},
        )

    def grant_verification_permission(self):
        permission = Permission.objects.get(
            codename="change_transportpayment",
            content_type__app_label="transport_pro",
            content_type__model="transportpayment",
        )
        self.role.permissions.add(permission)

    def test_unauthenticated_request_redirects_to_login(self):
        response = self.client.post(self.url)
        self.assertEqual(response.status_code, 302)
        self.payment.refresh_from_db()
        self.assertFalse(self.payment.verified)

    def test_staff_without_permission_is_forbidden(self):
        self.client.force_login(self.staff)
        response = self.client.post(self.url)
        self.assertEqual(response.status_code, 403)
        self.payment.refresh_from_db()
        self.assertFalse(self.payment.verified)
        self.assertEqual(TransportAuditEvent.objects.count(), 0)

    def test_authorized_user_verifies_payment_and_recalculates_balance(self):
        self.grant_verification_permission()
        self.client.force_login(self.staff)

        response = self.client.post(self.url)

        self.assertEqual(response.status_code, 302)
        self.payment.refresh_from_db()
        self.charge.refresh_from_db()
        self.assertTrue(self.payment.verified)
        self.assertEqual(self.payment.verified_by, self.staff)
        self.assertEqual(self.charge.status, "PARTIAL")
        self.assertEqual(
            TransportAuditEvent.objects.filter(
                action="VERIFY_PAYMENT",
                object_id=str(self.payment.pk),
            ).count(),
            1,
        )

    def test_repeated_verification_does_not_duplicate_audit_event(self):
        self.grant_verification_permission()
        self.client.force_login(self.staff)

        first = self.client.post(self.url)
        second = self.client.post(self.url)

        self.assertEqual(first.status_code, 302)
        self.assertEqual(second.status_code, 302)
        self.assertEqual(
            TransportAuditEvent.objects.filter(
                action="VERIFY_PAYMENT",
                object_id=str(self.payment.pk),
            ).count(),
            1,
        )

    def test_waived_charge_remains_waived(self):
        self.charge.status = "WAIVED"
        self.charge.save(update_fields=["status", "updated_at"])
        self.grant_verification_permission()
        self.client.force_login(self.staff)

        response = self.client.post(self.url)

        self.assertEqual(response.status_code, 302)
        self.charge.refresh_from_db()
        self.assertEqual(self.charge.status, "WAIVED")

    def test_get_request_cannot_verify_payment(self):
        self.grant_verification_permission()
        self.client.force_login(self.staff)

        response = self.client.get(self.url)

        self.assertEqual(response.status_code, 405)
        self.payment.refresh_from_db()
        self.assertFalse(self.payment.verified)


    def test_unauthorized_user_cannot_view_vehicle_list(self):
        self.client.force_login(self.staff)
        response = self.client.get(reverse(
            "transport_pro:list", kwargs={"kind": "vehicles"}
        ))
        self.assertEqual(response.status_code, 403)

    def test_unauthorized_user_cannot_create_payment(self):
        self.client.force_login(self.staff)
        response = self.client.get(reverse(
            "transport_pro:create", kwargs={"kind": "payments"}
        ))
        self.assertEqual(response.status_code, 403)
        self.assertEqual(TransportPayment.objects.count(), 1)

    def test_payment_cannot_be_edited_through_generic_route(self):
        self.grant_verification_permission()
        self.client.force_login(self.staff)
        response = self.client.get(reverse(
            "transport_pro:edit",
            kwargs={"kind": "payments", "pk": self.payment.pk},
        ))
        self.assertEqual(response.status_code, 403)

    def test_unauthorized_user_cannot_change_trip_status(self):
        from .models import Vehicle, TransportStaff, TransportTrip

        route = TransportRoute.objects.get(code="PAYTEST")
        vehicle = Vehicle.objects.create(
            registration_number="KTEST 701A", seating_capacity=14
        )
        driver = TransportStaff.objects.create(
            full_name="Trip Permission Test Driver", role="DRIVER"
        )
        trip = TransportTrip.objects.create(
            route=route, direction="AM", vehicle=vehicle, driver=driver
        )
        self.client.force_login(self.staff)
        response = self.client.post(reverse(
            "transport_pro:trip_action",
            kwargs={"pk": trip.pk, "action": "start"},
        ))
        self.assertEqual(response.status_code, 403)
        trip.refresh_from_db()
        self.assertEqual(trip.status, "PLANNED")

    def test_authorized_role_can_view_vehicle_list(self):
        permission = Permission.objects.get(
            codename="view_vehicle",
            content_type__app_label="transport_pro",
            content_type__model="vehicle",
        )
        self.role.permissions.add(permission)
        self.client.force_login(self.staff)
        response = self.client.get(reverse(
            "transport_pro:list", kwargs={"kind": "vehicles"}
        ))
        self.assertEqual(response.status_code, 200)

    def grant_summary_permissions(self):
        required = (
            ("view_vehicle", "vehicle"),
            ("view_transportstaff", "transportstaff"),
            ("view_learnertransport", "learnertransport"),
            ("view_transportroute", "transportroute"),
            ("view_transporttrip", "transporttrip"),
            ("view_transportincident", "transportincident"),
            ("view_maintenancerecord", "maintenancerecord"),
            ("view_boardingevent", "boardingevent"),
            ("view_transportcharge", "transportcharge"),
            ("view_transportpayment", "transportpayment"),
        )
        for codename, model_name in required:
            permission = Permission.objects.get(
                codename=codename,
                content_type__app_label="transport_pro",
                content_type__model=model_name,
            )
            self.role.permissions.add(permission)

    def test_dashboard_denies_trip_only_permission(self):
        permission = Permission.objects.get(
            codename="view_transporttrip",
            content_type__app_label="transport_pro",
            content_type__model="transporttrip",
        )
        self.role.permissions.add(permission)
        self.client.force_login(self.staff)
        response = self.client.get(reverse("transport_pro:dashboard"))
        self.assertEqual(response.status_code, 403)

    def test_reports_denies_trip_only_permission(self):
        permission = Permission.objects.get(
            codename="view_transporttrip",
            content_type__app_label="transport_pro",
            content_type__model="transporttrip",
        )
        self.role.permissions.add(permission)
        self.client.force_login(self.staff)
        response = self.client.get(reverse("transport_pro:reports"))
        self.assertEqual(response.status_code, 403)

    def test_authorized_role_can_view_dashboard_and_reports(self):
        self.grant_summary_permissions()
        self.client.force_login(self.staff)

        dashboard_response = self.client.get(reverse("transport_pro:dashboard"))
        reports_response = self.client.get(reverse("transport_pro:reports"))

        self.assertEqual(dashboard_response.status_code, 200)
        self.assertEqual(reports_response.status_code, 200)

    def test_outstanding_balance_uses_verified_payments_only(self):
        from .views import _outstanding_transport_balance

        self.assertEqual(
            _outstanding_transport_balance(), Decimal("1000.00")
        )

        self.payment.verified = True
        self.payment.save(update_fields=["verified", "updated_at"])

        self.assertEqual(
            _outstanding_transport_balance(), Decimal("400.00")
        )

    def test_waived_charge_is_excluded_from_outstanding_balance(self):
        from .views import _outstanding_transport_balance

        self.charge.status = "WAIVED"
        self.charge.save(update_fields=["status", "updated_at"])

        self.assertEqual(
            _outstanding_transport_balance(), Decimal("0.00")
        )

    def test_trip_status_is_not_exposed_in_generic_form(self):
        from .forms import TransportTripForm

        self.assertNotIn("status", TransportTripForm().fields)


from io import BytesIO
from datetime import date, time
from django.core.files.uploadedfile import SimpleUploadedFile
from openpyxl import Workbook, load_workbook
from students.models import Student



class LearnerTransportRegistrationTests(TestCase):
    """Regression tests for manual learner registration and Excel imports."""

    def setUp(self):
        self.user_model = get_user_model()
        self.staff = self.user_model.objects.create_user(
            username="registration_regression_staff",
            password="test-password-not-for-production",
        )
        self.role = Role.objects.create(
            name="Registration Regression Role",
            is_active=True,
        )
        UserProfile.objects.update_or_create(
            user=self.staff,
            defaults={"role": "STAFF", "custom_role": self.role},
        )

        self.route = TransportRoute.objects.create(
            name="Registration Test Route",
            code="RTR1",
            area="Test Area",
            active=True,
        )
        self.pickup = RouteStop.objects.create(
            route=self.route,
            name="Pickup Point",
            sequence=1,
            morning_pickup_time=time(7, 30),
            afternoon_dropoff_time=time(16, 0),
            active=True,
        )
        self.dropoff = RouteStop.objects.create(
            route=self.route,
            name="School Gate",
            sequence=2,
            morning_pickup_time=time(7, 40),
            afternoon_dropoff_time=time(16, 15),
            active=True,
        )

        self.student = Student.objects.create(
            admission_no="REG001",
            first_name="Amina",
            last_name="Test",
            gender="Female",
            date_of_birth=date(2015, 1, 1),
            class_name="Grade 4",
            stream="Blue",
            parent_name="Parent Test",
            parent_phone="0712345678",
            is_active=True,
        )

        self.registration_url = reverse(
            "transport_pro:learner_registration"
        )
        self.template_url = reverse(
            "transport_pro:learner_import_template"
        )
        self.import_url = reverse(
            "transport_pro:learner_bulk_import"
        )

    def grant_registration_permissions(self, *, assignment=True, charge=True):
        permission_specs = []
        if assignment:
            permission_specs.append(
                ("add_learnertransport", "learnertransport")
            )
        if charge:
            permission_specs.append(
                ("add_transportcharge", "transportcharge")
            )

        for codename, model_name in permission_specs:
            permission = Permission.objects.get(
                codename=codename,
                content_type__app_label="transport_pro",
                content_type__model=model_name,
            )
            self.role.permissions.add(permission)

    def workbook_upload(self, rows):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append([
            "Admission Number",
            "Pickup Area / Stop",
            "Pickup Time (HH:MM)",
            "Drop-off Area / Stop",
            "Drop-off Time (HH:MM)",
            "Price to Charge",
        ])
        for row in rows:
            sheet.append(row)

        buffer = BytesIO()
        workbook.save(buffer)
        workbook.close()

        return SimpleUploadedFile(
            "learners.xlsx",
            buffer.getvalue(),
            content_type=(
                "application/vnd.openxmlformats-officedocument."
                "spreadsheetml.sheet"
            ),
        )

    def valid_import_row(self, admission_no="REG001"):
        return [
            admission_no,
            f"Test Area | RTR1 - {self.pickup.name}",
            "07:30",
            f"Test Area | RTR1 - {self.dropoff.name}",
            "16:15",
            2500,
        ]

    def test_registration_requires_both_permissions(self):
        self.client.force_login(self.staff)

        # Assignment permission alone must not authorize a workflow
        # that also creates a transport charge.
        self.grant_registration_permissions(assignment=True, charge=False)
        response = self.client.post(self.registration_url, {
            "admission_no": "REG001",
            "pickup_stop": str(self.pickup.pk),
            "dropoff_stop": str(self.dropoff.pk),
            "pickup_time": "07:30",
            "dropoff_time": "16:15",
            "period_label": "Term 3 2026",
            "price": "2500",
        })

        self.assertEqual(response.status_code, 403)
        self.assertEqual(LearnerTransport.objects.count(), 0)
        self.assertEqual(TransportCharge.objects.count(), 0)

    def test_manual_registration_creates_assignment_and_due_charge(self):
        self.grant_registration_permissions()
        self.client.force_login(self.staff)

        response = self.client.post(self.registration_url, {
            "admission_no": "REG001",
            "pickup_stop": str(self.pickup.pk),
            "dropoff_stop": str(self.dropoff.pk),
            "pickup_time": "07:30",
            "dropoff_time": "16:15",
            "period_label": "Term 3 2026",
            "price": "2500",
        })

        self.assertIn(response.status_code, (200, 302))
        assignment = LearnerTransport.objects.get(
            learner_admission_no__iexact="reg001"
        )
        self.assertEqual(assignment.learner_name, "Amina Test")
        self.assertEqual(assignment.route_id, self.route.pk)
        self.assertEqual(assignment.pickup_stop_id, self.pickup.pk)
        self.assertEqual(assignment.dropoff_stop_id, self.dropoff.pk)
        self.assertEqual(assignment.status, "PENDING")

        charge = TransportCharge.objects.get(learner=assignment)
        self.assertEqual(charge.period_label, "Term 3 2026")
        self.assertEqual(charge.amount, Decimal("2500"))
        self.assertEqual(charge.status, "DUE")

    def test_manual_registration_rejects_case_insensitive_duplicate(self):
        self.grant_registration_permissions()
        self.client.force_login(self.staff)

        LearnerTransport.objects.create(
            learner_admission_no="reg001",
            learner_name="Existing Assignment",
            class_name="Grade 4",
            parent_name="Parent Test",
            parent_phone="0712345678",
            route=self.route,
            stop=self.pickup,
            pickup_stop=self.pickup,
            dropoff_stop=self.dropoff,
            direction="BOTH",
            status="ACTIVE",
        )

        response = self.client.post(self.registration_url, {
            "admission_no": "REG001",
            "pickup_stop": str(self.pickup.pk),
            "dropoff_stop": str(self.dropoff.pk),
            "pickup_time": "07:30",
            "dropoff_time": "16:15",
            "period_label": "Term 3 2026",
            "price": "2500",
        })

        self.assertEqual(LearnerTransport.objects.count(), 1)
        self.assertEqual(TransportCharge.objects.count(), 0)
        self.assertIn(response.status_code, (200, 302))

    def test_excel_import_creates_assignment_and_due_charge(self):
        self.grant_registration_permissions()
        self.client.force_login(self.staff)

        response = self.client.post(self.import_url, {
            "workbook": self.workbook_upload([
                self.valid_import_row("REG001"),
            ]),
            "period_label": "Term 3 2026",
        })

        self.assertIn(response.status_code, (200, 302))
        self.assertEqual(LearnerTransport.objects.count(), 1)
        self.assertEqual(TransportCharge.objects.count(), 1)

        assignment = LearnerTransport.objects.get()
        self.assertEqual(assignment.learner_admission_no, "REG001")
        self.assertEqual(assignment.pickup_stop_id, self.pickup.pk)
        self.assertEqual(assignment.dropoff_stop_id, self.dropoff.pk)

        charge = TransportCharge.objects.get(learner=assignment)
        self.assertEqual(charge.amount, Decimal("2500"))
        self.assertEqual(charge.status, "DUE")
        self.assertEqual(charge.period_label, timezone.localdate().strftime("%B %Y"))

    def test_invalid_excel_batch_saves_no_assignments_or_charges(self):
        self.grant_registration_permissions()
        self.client.force_login(self.staff)

        response = self.client.post(self.import_url, {
            "workbook": self.workbook_upload([
                self.valid_import_row("REG001"),
                self.valid_import_row("UNKNOWN999"),
            ]),
            "period_label": "Term 3 2026",
        })

        self.assertIn(response.status_code, (200, 302))
        self.assertEqual(LearnerTransport.objects.count(), 0)
        self.assertEqual(TransportCharge.objects.count(), 0)

    def test_excel_import_rejects_duplicate_admissions_ignoring_case(self):
        self.grant_registration_permissions()
        self.client.force_login(self.staff)

        response = self.client.post(self.import_url, {
            "workbook": self.workbook_upload([
                self.valid_import_row("REG001"),
                self.valid_import_row("reg001"),
            ]),
            "period_label": "Term 3 2026",
        })

        self.assertIn(response.status_code, (200, 302))
        self.assertEqual(LearnerTransport.objects.count(), 0)
        self.assertEqual(TransportCharge.objects.count(), 0)

    def test_excel_template_download_is_a_valid_workbook(self):
        self.grant_registration_permissions()
        self.client.force_login(self.staff)

        response = self.client.get(self.template_url)

        self.assertEqual(response.status_code, 200)
        self.assertIn(
            "spreadsheetml.sheet",
            response.get("Content-Type", ""),
        )
        workbook = load_workbook(BytesIO(response.content), read_only=True)
        sheet = workbook.active
        headers = [cell.value for cell in sheet[1]]
        self.assertEqual(headers, [
            "Admission Number",
            "Pickup Area / Stop",
            "Pickup Time (HH:MM)",
            "Drop-off Area / Stop",
            "Drop-off Time (HH:MM)",
            "Price to Charge",
        ])
        workbook.close()

