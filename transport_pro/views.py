from decimal import Decimal
from django.contrib import messages
from django.contrib.auth.decorators import login_required, permission_required
from django.core.exceptions import ValidationError
from django.db import transaction
from django.db.models import Sum, Count, Q
from django.http import HttpResponse, HttpResponseForbidden, JsonResponse
from django.shortcuts import render, redirect, get_object_or_404
from django.utils import timezone
from django.views.decorators.http import require_POST
from . import models, forms


def _audit(request, action, obj, details=None):
    models.TransportAuditEvent.objects.create(
        actor=request.user if request.user.is_authenticated else None,
        action=action, object_type=obj.__class__.__name__, object_id=str(obj.pk),
        details=details or {},
        ip_address=(request.META.get("HTTP_X_FORWARDED_FOR", "").split(",")[0].strip()
                    or request.META.get("REMOTE_ADDR") or None)
    )


def _has_transport_permission(user, codename, model_name):
    """Fail-closed authorization through the active custom ERP role."""
    if not user.is_authenticated:
        return False
    if user.is_superuser:
        return True
    try:
        role = user.profile.custom_role
    except Exception:
        return False
    if not role or not role.is_active:
        return False
    return role.permissions.filter(
        codename=codename,
        content_type__app_label="transport_pro",
        content_type__model=model_name,
    ).exists()


def _can_verify_transport_payments(user):
    return _has_transport_permission(
        user, "change_transportpayment", "transportpayment"
    )


def _has_transport_summary_permissions(user, required_permissions):
    """Allow access only when every permission for a page's data is present."""
    return all(
        _has_transport_permission(user, codename, model_name)
        for codename, model_name in required_permissions
    )


def _has_transport_dashboard_access(user):
    return _has_transport_summary_permissions(user, (
        ("view_vehicle", "vehicle"),
        ("view_transportstaff", "transportstaff"),
        ("view_learnertransport", "learnertransport"),
        ("view_transportroute", "transportroute"),
        ("view_transporttrip", "transporttrip"),
        ("view_transportincident", "transportincident"),
        ("view_maintenancerecord", "maintenancerecord"),
        ("view_boardingevent", "boardingevent"),
        ("view_transportcharge", "transportcharge"),
        ("view_transportpayment", "transportpayment"),
    ))


def _has_transport_reports_access(user):
    return _has_transport_summary_permissions(user, (
        ("view_learnertransport", "learnertransport"),
        ("view_transportroute", "transportroute"),
        ("view_transporttrip", "transporttrip"),
        ("view_transportincident", "transportincident"),
        ("view_maintenancerecord", "maintenancerecord"),
        ("view_transportcharge", "transportcharge"),
        ("view_transportpayment", "transportpayment"),
    ))


def _outstanding_transport_balance():
    """Sum each charge's positive balance after verified payments."""
    from django.db.models import Case, When, F, Value, DecimalField
    from django.db.models.functions import Coalesce

    money = DecimalField(max_digits=14, decimal_places=2)
    charges = models.TransportCharge.objects.filter(
        status__in=["DUE", "PARTIAL"]
    ).annotate(
        verified_paid=Coalesce(
            Sum("payments__amount", filter=Q(payments__verified=True)),
            Value(Decimal("0.00")),
            output_field=money,
        )
    ).annotate(
        raw_balance=F("amount") - F("verified_paid")
    )
    total = charges.aggregate(
        total=Sum(Case(
            When(raw_balance__gt=0, then=F("raw_balance")),
            default=Value(Decimal("0.00")),
            output_field=money,
        ))
    )["total"]
    return total or Decimal("0.00")




@login_required
def dashboard(request):
    if not _has_transport_dashboard_access(request.user):
        return HttpResponseForbidden(
            "You do not have permission to view all transport dashboard summaries."
        )
    today = timezone.localdate()
    trips = models.TransportTrip.objects.filter(service_date=today).select_related("route","vehicle","driver")
    context = {
        "today": today,
        "vehicles_total": models.Vehicle.objects.count(),
        "vehicles_serviceable": models.Vehicle.objects.filter(status="AVAILABLE").count(),
        "active_learners": models.LearnerTransport.objects.filter(status="ACTIVE").count(),
        "active_routes": models.TransportRoute.objects.filter(active=True).count(),
        "today_trips": trips,
        "open_incidents": models.TransportIncident.objects.exclude(status="RESOLVED").count(),
        "open_maintenance": models.MaintenanceRecord.objects.exclude(status="COMPLETED").count(),
        "document_alerts": sum(1 for v in models.Vehicle.objects.all() if v.documents_expired),
        "today_boarding": models.BoardingEvent.objects.filter(trip__service_date=today).count(),
        "outstanding": _outstanding_transport_balance(),
    }
    return render(request, "transport_pro/dashboard.html", context)


def _learner_registration_allowed(user):
    """Registration creates both an assignment and a charge; require both rights."""
    return (
        _has_transport_permission(user, "add_learnertransport", "learnertransport")
        and _has_transport_permission(user, "add_transportcharge", "transportcharge")
    )


def _active_registration_stops():
    return list(
        models.RouteStop.objects.filter(
            active=True, route__active=True
        ).select_related("route").order_by(
            "route__area", "route__name", "sequence", "name"
        )
    )


def _registration_stop_label(stop):
    area = (stop.route.area or stop.route.name).strip()
    return f"{area} | {stop.route.code} - {stop.name}"


def _student_registration_data(admission_no):
    from students.models import Student

    student = Student.objects.filter(
        admission_no__iexact=admission_no.strip(), is_active=True
    ).first()
    if not student:
        return None

    full_name = " ".join(
        part.strip() for part in (
            student.first_name or "",
            student.middle_name or "",
            student.last_name or "",
        ) if part and part.strip()
    )
    parent_name = (student.parent_name or "").strip()
    parent_phone = (student.parent_phone or "").strip()

    if not parent_name:
        parent = student.parents.select_related("user").first()
        if parent:
            parent_name = parent.full_name or ""
            parent_phone = parent_phone or parent.phone or ""

    return {
        "admission_no": student.admission_no,
        "learner_name": full_name,
        "class_name": student.class_name or "",
        "parent_name": parent_name,
        "parent_phone": parent_phone,
    }


def _registration_context():
    stops = _active_registration_stops()
    stop_data = []
    for stop in stops:
        stop_data.append({
            "id": stop.pk,
            "label": _registration_stop_label(stop),
            "route_id": stop.route_id,
            "morning_time": (
                stop.morning_pickup_time.strftime("%H:%M")
                if stop.morning_pickup_time else ""
            ),
            "afternoon_time": (
                stop.afternoon_dropoff_time.strftime("%H:%M")
                if stop.afternoon_dropoff_time else ""
            ),
        })
    return {
        "stops": stops,
        "stop_data": stop_data,
        "default_period": timezone.localdate().strftime("%B %Y"),
    }


@login_required
def learner_lookup(request):
    if not _learner_registration_allowed(request.user):
        return HttpResponseForbidden("You do not have permission to register learners and create transport charges.")
    admission_no = request.GET.get("admission_no", "").strip()
    if not admission_no:
        return JsonResponse({"ok": False, "error": "Enter an admission number."}, status=400)
    data = _student_registration_data(admission_no)
    if not data:
        return JsonResponse({
            "ok": False,
            "error": "No active learner was found with that admission number."
        }, status=404)
    return JsonResponse({"ok": True, "learner": data})


@login_required
def learner_registration(request):
    from decimal import InvalidOperation
    from students.models import Student

    if not _learner_registration_allowed(request.user):
        return HttpResponseForbidden("You do not have permission to register learners and create transport charges.")

    context = _registration_context()
    errors = []

    if request.method == "POST":
        admission_no = request.POST.get("admission_no", "").strip()
        data = _student_registration_data(admission_no)
        pickup_id = request.POST.get("pickup_stop", "").strip()
        dropoff_id = request.POST.get("dropoff_stop", "").strip()
        pickup_time = request.POST.get("pickup_time", "").strip()
        dropoff_time = request.POST.get("dropoff_time", "").strip()
        period = request.POST.get("period_label", "").strip() or context["default_period"]
        price_text = request.POST.get("price", "").strip()

        if not data:
            errors.append("Enter a valid admission number for an active learner.")
        try:
            pickup = models.RouteStop.objects.select_related("route").get(
                pk=int(pickup_id), active=True, route__active=True
            )
        except (ValueError, TypeError, models.RouteStop.DoesNotExist):
            pickup = None
            errors.append("Select an active pickup area.")
        try:
            dropoff = models.RouteStop.objects.select_related("route").get(
                pk=int(dropoff_id), active=True, route__active=True
            )
        except (ValueError, TypeError, models.RouteStop.DoesNotExist):
            dropoff = None
            errors.append("Select an active drop-off area.")

        if pickup and (
            not pickup.morning_pickup_time
            or pickup_time != pickup.morning_pickup_time.strftime("%H:%M")
        ):
            errors.append("The pickup time must match the configured time for the selected pickup stop.")
        if dropoff and (
            not dropoff.afternoon_dropoff_time
            or dropoff_time != dropoff.afternoon_dropoff_time.strftime("%H:%M")
        ):
            errors.append("The drop-off time must match the configured time for the selected drop-off stop.")

        try:
            price = Decimal(price_text)
            if not price.is_finite() or price < 0:
                raise InvalidOperation
        except (InvalidOperation, ValueError):
            price = None
            errors.append("Enter a valid, non-negative transport charge.")

        if not period or len(period) > 80:
            errors.append("Enter a billing period of no more than 80 characters.")

        if data and models.LearnerTransport.objects.filter(
            learner_admission_no__iexact=data["admission_no"],
            status__in=["PENDING", "ACTIVE"],
        ).exists():
            errors.append("This learner already has a pending or active transport assignment.")

        if not errors:
            try:
                with transaction.atomic():
                    student = Student.objects.get(
                        admission_no__iexact=data["admission_no"], is_active=True
                    )
                    assignment = models.LearnerTransport(
                        learner_admission_no=student.admission_no,
                        learner_name=data["learner_name"],
                        class_name=data["class_name"],
                        parent_name=data["parent_name"],
                        parent_phone=data["parent_phone"],
                        route=pickup.route,
                        stop=pickup,  # legacy compatibility
                        pickup_stop=pickup,
                        dropoff_stop=dropoff,
                        pickup_time=pickup.morning_pickup_time,
                        dropoff_time=dropoff.afternoon_dropoff_time,
                        direction="BOTH",
                        status="PENDING",
                    )
                    assignment.full_clean()
                    assignment.save()
                    charge = models.TransportCharge.objects.create(
                        learner=assignment,
                        period_label=period,
                        amount=price,
                        status="DUE",
                        notes="Created during learner transport registration.",
                    )
                    _audit(request, "CREATE", assignment, {
                        "source": "learner_registration",
                        "charge_id": charge.pk,
                        "pickup_stop_id": pickup.pk,
                        "dropoff_stop_id": dropoff.pk,
                    })
                    _audit(request, "CREATE", charge, {
                        "source": "learner_registration",
                        "learner_assignment_id": assignment.pk,
                    })
                messages.success(request, f"Transport registration saved for {data['learner_name']}. Charge is due.")
                return redirect("transport_pro:learner_registration")
            except ValidationError as exc:
                errors.extend(exc.messages)
            except Exception:
                # Atomic block rolls back both the assignment and charge.
                errors.append("The registration could not be saved. No changes were committed; check the details and try again.")

    context.update({"errors": errors})
    return render(request, "transport_pro/learner_registration.html", context)


@login_required
def learner_import_template(request):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.workbook.defined_name import DefinedName

    if not _learner_registration_allowed(request.user):
        return HttpResponseForbidden("You do not have permission to import learner registrations.")

    stops = _active_registration_stops()
    pickups = sorted({
        _registration_stop_label(s) for s in stops if s.morning_pickup_time
    })
    dropoffs = sorted({
        _registration_stop_label(s) for s in stops if s.afternoon_dropoff_time
    })
    pickup_times = sorted({
        s.morning_pickup_time.strftime("%H:%M")
        for s in stops if s.morning_pickup_time
    })
    dropoff_times = sorted({
        s.afternoon_dropoff_time.strftime("%H:%M")
        for s in stops if s.afternoon_dropoff_time
    })

    if not pickups or not dropoffs or not pickup_times or not dropoff_times:
        return HttpResponse(
            "Configure active routes and stops with morning pickup and afternoon drop-off times before downloading the template.",
            status=400,
        )

    wb = Workbook()
    ws = wb.active
    ws.title = "Learner Registrations"
    headers = [
        "Admission Number", "Pickup Area / Stop", "Pickup Time (HH:MM)",
        "Drop-off Area / Stop", "Drop-off Time (HH:MM)", "Price to Charge",
    ]
    ws.append(headers)
    ws.append(["ADM001", pickups[0], pickup_times[0], dropoffs[0], dropoff_times[0], 2500])

    navy = "17365D"
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=navy)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 34
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = "A1:F501"
    widths = [23, 38, 23, 38, 25, 20]
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[chr(64 + index)].width = width

    ws["A3"] = "Instructions"
    ws["A3"].font = Font(bold=True, color=navy)
    ws["A4"] = "Replace the example row. Use the dropdowns for areas and times."
    ws["A5"] = "Admission numbers must belong to active learners in the school register."
    ws["A6"] = "Do not change the column headings. A failed row prevents the entire batch from saving."
    ws["A7"] = "The price creates a Due charge; it does not record a payment."
    ws.merge_cells("A4:F4")
    ws.merge_cells("A5:F5")
    ws.merge_cells("A6:F6")
    ws.merge_cells("A7:F7")

    options = wb.create_sheet("Options")
    for col, (name, values) in enumerate((
        ("PickupAreas", pickups),
        ("PickupTimes", pickup_times),
        ("DropoffAreas", dropoffs),
        ("DropoffTimes", dropoff_times),
    ), start=1):
        options.cell(1, col, name)
        for row, value in enumerate(values, start=2):
            options.cell(row, col, value)
        letter = chr(64 + col)
        wb.defined_names.add(DefinedName(
            name, attr_text=f"'Options'!${letter}$2:${letter}${len(values)+1}"
        ))
    options.sheet_state = "hidden"

    validations = [
        ("B2:B501", "=PickupAreas", "Choose an active pickup stop."),
        ("C2:C501", "=PickupTimes", "Choose a configured pickup time."),
        ("D2:D501", "=DropoffAreas", "Choose an active drop-off stop."),
        ("E2:E501", "=DropoffTimes", "Choose a configured drop-off time."),
    ]
    for cell_range, formula, prompt in validations:
        dv = DataValidation(type="list", formula1=formula, allow_blank=False)
        dv.errorTitle = "Invalid selection"
        dv.error = "Select a value from the dropdown list."
        dv.promptTitle = "Configured transport options"
        dv.prompt = prompt
        dv.showErrorMessage = True
        dv.showInputMessage = True
        ws.add_data_validation(dv)
        dv.add(cell_range)

    price_validation = DataValidation(
        type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=False
    )
    price_validation.showErrorMessage = True
    price_validation.error = "Enter a non-negative amount."
    ws.add_data_validation(price_validation)
    price_validation.add("F2:F501")

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = 'attachment; filename="transport_learner_import_template.xlsx"'
    wb.save(response)
    return response


@login_required
def learner_bulk_import(request):
    from datetime import datetime, time
    from decimal import InvalidOperation
    from io import BytesIO
    from openpyxl import load_workbook
    from openpyxl.utils.datetime import from_excel
    from students.models import Student

    if not _learner_registration_allowed(request.user):
        return HttpResponseForbidden("You do not have permission to import learner registrations.")

    context = _registration_context()
    errors = []
    imported = request.GET.get("imported", "")
    if request.method == "POST":
        uploaded = request.FILES.get("workbook")
        if not uploaded or not uploaded.name.lower().endswith(".xlsx"):
            errors.append("Upload an Excel .xlsx workbook.")
        elif uploaded.size > 5 * 1024 * 1024:
            errors.append("The workbook must be 5 MB or smaller.")
        else:
            try:
                wb = load_workbook(BytesIO(uploaded.read()), read_only=True, data_only=True)
                ws = wb["Learner Registrations"] if "Learner Registrations" in wb.sheetnames else wb.active
                expected = [
                    "Admission Number", "Pickup Area / Stop", "Pickup Time (HH:MM)",
                    "Drop-off Area / Stop", "Drop-off Time (HH:MM)", "Price to Charge",
                ]
                actual = [str(ws.cell(1, col).value or "").strip() for col in range(1, 7)]
                if actual != expected:
                    errors.append("The workbook headings do not match the official template. Download a fresh template.")
                elif ws.max_row > 502:
                    errors.append("The workbook has too many rows. Import no more than 500 learners at a time.")
                else:
                    stops = _active_registration_stops()
                    pickup_map = {
                        _registration_stop_label(s): s for s in stops if s.morning_pickup_time
                    }
                    dropoff_map = {
                        _registration_stop_label(s): s for s in stops if s.afternoon_dropoff_time
                    }
                    seen = set()
                    prepared = []
                    for row_no in range(2, ws.max_row + 1):
                        values = [ws.cell(row_no, col).value for col in range(1, 7)]
                        if all(v is None or str(v).strip() == "" for v in values):
                            continue
                        admission = str(values[0] or "").strip()
                        pickup_label = str(values[1] or "").strip()
                        dropoff_label = str(values[3] or "").strip()

                        def time_text(value):
                            if isinstance(value, datetime):
                                return value.strftime("%H:%M")
                            if isinstance(value, time):
                                return value.strftime("%H:%M")
                            if isinstance(value, (int, float)):
                                try:
                                    converted = from_excel(value)
                                    if isinstance(converted, datetime):
                                        return converted.strftime("%H:%M")
                                    if isinstance(converted, time):
                                        return converted.strftime("%H:%M")
                                except Exception:
                                    pass
                            raw = str(value or "").strip()
                            for fmt in ("%H:%M", "%I:%M %p"):
                                try:
                                    return datetime.strptime(raw, fmt).strftime("%H:%M")
                                except ValueError:
                                    pass
                            return ""

                        pickup_time_text = time_text(values[2])
                        dropoff_time_text = time_text(values[4])
                        row_errors = []

                        if not admission:
                            row_errors.append("missing admission number")
                        key = admission.casefold()
                        if key and key in seen:
                            row_errors.append("duplicate admission number in this workbook")
                        if key:
                            seen.add(key)

                        learner = _student_registration_data(admission) if admission else None
                        if admission and not learner:
                            row_errors.append("learner not found or inactive")
                        pickup = pickup_map.get(pickup_label)
                        dropoff = dropoff_map.get(dropoff_label)
                        if not pickup:
                            row_errors.append("pickup area/stop is invalid or inactive")
                        if not dropoff:
                            row_errors.append("drop-off area/stop is invalid or inactive")
                        if pickup and pickup_time_text != pickup.morning_pickup_time.strftime("%H:%M"):
                            row_errors.append("pickup time does not match the selected stop")
                        if dropoff and dropoff_time_text != dropoff.afternoon_dropoff_time.strftime("%H:%M"):
                            row_errors.append("drop-off time does not match the selected stop")
                        try:
                            amount = Decimal(str(values[5]).strip())
                            if not amount.is_finite() or amount < 0:
                                raise InvalidOperation
                        except (InvalidOperation, ValueError, TypeError):
                            amount = None
                            row_errors.append("price must be a non-negative number")

                        if learner and models.LearnerTransport.objects.filter(
                            learner_admission_no__iexact=learner["admission_no"],
                            status__in=["PENDING", "ACTIVE"],
                        ).exists():
                            row_errors.append("learner already has a pending or active assignment")

                        if row_errors:
                            errors.append(f"Row {row_no}: " + "; ".join(row_errors))
                        else:
                            prepared.append({
                                "learner": learner, "pickup": pickup, "dropoff": dropoff,
                                "pickup_time": pickup.morning_pickup_time,
                                "dropoff_time": dropoff.afternoon_dropoff_time,
                                "amount": amount, "row_no": row_no,
                            })
                    wb.close()

                    if not prepared and not errors:
                        errors.append("The workbook contains no learner rows to import.")

                    if not errors:
                        try:
                            with transaction.atomic():
                                # Recheck duplicates inside the transaction to reduce race conditions.
                                admission_numbers = [p["learner"]["admission_no"] for p in prepared]
                                # Use case-insensitive checks inside the transaction as well.
                                # A case-sensitive __in query can miss an existing assignment
                                # when admission-number casing differs.
                                duplicate_exists = any(
                                    models.LearnerTransport.objects.filter(
                                        learner_admission_no__iexact=admission,
                                        status__in=["PENDING", "ACTIVE"],
                                    ).exists()
                                    for admission in admission_numbers
                                )
                                if duplicate_exists:
                                    raise ValueError("One or more learners already have pending or active assignments. Refresh and validate the workbook again.")

                                period = timezone.localdate().strftime("%B %Y")
                                for item in prepared:
                                    learner = item["learner"]
                                    assignment = models.LearnerTransport(
                                        learner_admission_no=learner["admission_no"],
                                        learner_name=learner["learner_name"],
                                        class_name=learner["class_name"],
                                        parent_name=learner["parent_name"],
                                        parent_phone=learner["parent_phone"],
                                        route=item["pickup"].route,
                                        stop=item["pickup"],
                                        pickup_stop=item["pickup"],
                                        dropoff_stop=item["dropoff"],
                                        pickup_time=item["pickup_time"],
                                        dropoff_time=item["dropoff_time"],
                                        direction="BOTH",
                                        status="PENDING",
                                    )
                                    assignment.full_clean()
                                    assignment.save()
                                    charge = models.TransportCharge.objects.create(
                                        learner=assignment, period_label=period,
                                        amount=item["amount"], status="DUE",
                                        notes="Created by Excel learner transport import.",
                                    )
                                    _audit(request, "BULK_IMPORT", assignment, {
                                        "source": "xlsx",
                                        "row": item["row_no"],
                                        "charge_id": charge.pk,
                                    })
                            messages.success(request, f"Successfully imported {len(prepared)} learner registration(s). Charges are marked Due.")
                            return redirect(f"{request.path}?imported={len(prepared)}")
                        except Exception as exc:
                            errors.append(str(exc) if isinstance(exc, ValueError) else "The batch could not be saved. The transaction was rolled back; no partial import was committed.")
            except Exception:
                errors.append("The workbook could not be read. Use the latest downloaded .xlsx template and try again.")

    context.update({"errors": errors, "imported": imported})
    return render(request, "transport_pro/learner_bulk_import.html", context)


@login_required
def object_list(request, kind):
    config = {
        "vehicles": (models.Vehicle, "Vehicles"),
        "staff": (models.TransportStaff, "Drivers & Attendants"),
        "routes": (models.TransportRoute, "Routes"),
        "stops": (models.RouteStop, "Route Stops"),
        "learners": (models.LearnerTransport, "Learner Assignments"),
        "trips": (models.TransportTrip, "Trips"),
        "inspections": (models.VehicleInspection, "Vehicle Inspections"),
        "boarding": (models.BoardingEvent, "Boarding & Drop-off Events"),
        "maintenance": (models.MaintenanceRecord, "Maintenance"),
        "incidents": (models.TransportIncident, "Incidents"),
        "charges": (models.TransportCharge, "Transport Charges"),
        "payments": (models.TransportPayment, "Transport Payments"),
        "notifications": (models.TransportNotificationLog, "Notification Log"),
        "audit": (models.TransportAuditEvent, "Audit Trail"),
    }
    if kind not in config:
        return redirect("transport_pro:dashboard")
    model, title = config[kind]
    model_name = model._meta.model_name
    if not _has_transport_permission(request.user, f"view_{model_name}", model_name):
        return HttpResponseForbidden("You do not have permission to view this transport record type.")
    qs = model.objects.all()
    query = request.GET.get("q", "").strip()
    if query:
        fields = [f.name for f in model._meta.fields if getattr(f, "get_internal_type", lambda: "")() in ("CharField","TextField")]
        if fields:
            predicate = Q()
            for field in fields:
                predicate |= Q(**{field + "__icontains": query})
            qs = qs.filter(predicate)
    return render(request, "transport_pro/object_list.html", {"items": qs[:500], "title": title, "kind": kind, "query": query})


FORM_MAP = {
    "vehicles": (forms.VehicleForm, "Vehicle"),
    "staff": (forms.TransportStaffForm, "Transport staff"),
    "routes": (forms.TransportRouteForm, "Route"),
    "stops": (forms.RouteStopForm, "Route stop"),
    "learners": (forms.LearnerTransportForm, "Learner assignment"),
    "trips": (forms.TransportTripForm, "Trip"),
    "inspections": (forms.VehicleInspectionForm, "Vehicle inspection"),
    "boarding": (forms.BoardingEventForm, "Boarding event"),
    "maintenance": (forms.MaintenanceRecordForm, "Maintenance record"),
    "incidents": (forms.TransportIncidentForm, "Incident"),
    "charges": (forms.TransportChargeForm, "Transport charge"),
    "payments": (forms.TransportPaymentForm, "Transport payment"),
}


@login_required
def transport_employee_lookup(request):
    """Return eligible support-employee details for the staff registration form."""
    if not _has_transport_permission(
        request.user, "add_transportstaff", "transportstaff"
    ):
        return JsonResponse(
            {"ok": False, "message": "You do not have permission to register transport staff."},
            status=403,
        )

    number = (request.GET.get("employee_number") or "").strip()
    if not number:
        return JsonResponse({"ok": False, "message": "Enter an employee number."}, status=400)

    from hr_payroll.models import Employee
    from .forms import is_eligible_transport_employee

    try:
        employee = Employee.objects.select_related("department", "position").get(
            employee_number__iexact=number
        )
    except Employee.DoesNotExist:
        return JsonResponse(
            {"ok": False, "message": "No HR employee found with that employee number."},
            status=404,
        )

    if not is_eligible_transport_employee(employee):
        return JsonResponse({
            "ok": False,
            "message": (
                "Not eligible: the employee must be active support staff, not teaching "
                "or leadership staff, and must have a clear HR department or support position."
            ),
        }, status=400)

    if models.TransportStaff.objects.filter(employee=employee).exists():
        return JsonResponse(
            {"ok": False, "message": "This employee is already registered as transport staff."},
            status=409,
        )

    return JsonResponse({
        "ok": True,
        "employee_number": employee.employee_number,
        "full_name": " ".join(filter(None, [
            employee.first_name, employee.middle_name, employee.last_name
        ])),
        "phone": employee.phone or "",
        "emergency_contact": " ".join(filter(None, [
            employee.emergency_contact_name, employee.emergency_contact_phone
        ])),
        "department": employee.department.name if employee.department_id else "",
        "position": employee.position.title if employee.position_id else "",
    })

@login_required
def create_object(request, kind):
    if kind not in FORM_MAP:
        return redirect("transport_pro:dashboard")
    form_class, label = FORM_MAP[kind]
    model_name = form_class._meta.model._meta.model_name
    if not _has_transport_permission(request.user, f"add_{model_name}", model_name):
        return HttpResponseForbidden("You do not have permission to create this transport record type.")
    form = form_class(request.POST or None)
    if request.method == "POST" and form.is_valid():
        obj = form.save(commit=False)
        if isinstance(obj, models.TransportTrip):
            if obj.vehicle.documents_expired:
                form.add_error(None, "This vehicle has expired compliance documents; trip cannot be authorised.")
            elif obj.vehicle.status in ("MAINTENANCE","OUT_OF_SERVICE"):
                form.add_error(None, "This vehicle is not available for service.")
            elif not obj.driver.active or (obj.driver.licence_expiry and obj.driver.licence_expiry < timezone.localdate()):
                form.add_error(None, "Driver is inactive or has an expired licence.")
            else:
                obj.save()
                form.save_m2m()
                _audit(request, "CREATE", obj)
                messages.success(request, f"{label} created.")
                return redirect("transport_pro:list", kind=kind)
        else:
            obj.save()
            form.save_m2m()
            _audit(request, "CREATE", obj)
            messages.success(request, f"{label} created.")
            return redirect("transport_pro:list", kind=kind)
    return render(request, "transport_pro/object_form.html", {"form": form, "title": "Add " + label, "kind": kind})


@login_required
def edit_object(request, kind, pk):
    if kind not in FORM_MAP:
        return redirect("transport_pro:dashboard")
    form_class, label = FORM_MAP[kind]
    model = form_class._meta.model
    model_name = model._meta.model_name
    if kind == "payments":
        return HttpResponseForbidden("Recorded payments cannot be edited here. Use an approved correction process.")
    if not _has_transport_permission(request.user, f"change_{model_name}", model_name):
        return HttpResponseForbidden("You do not have permission to edit this transport record type.")
    obj = get_object_or_404(model, pk=pk)
    form = form_class(request.POST or None, instance=obj)
    if request.method == "POST" and form.is_valid():
        saved = form.save(commit=False)
        if isinstance(saved, models.TransportTrip):
            if saved.vehicle.documents_expired:
                form.add_error(None, "Vehicle has expired compliance documents; trip cannot be saved.")
            elif saved.vehicle.status in ("MAINTENANCE", "OUT_OF_SERVICE"):
                form.add_error(None, "Vehicle is not available for service.")
            elif not saved.driver.active or (saved.driver.licence_expiry and saved.driver.licence_expiry < timezone.localdate()):
                form.add_error(None, "Driver is inactive or has an expired licence.")
            if form.errors:
                return render(request, "transport_pro/object_form.html", {"form": form, "title": "Edit " + label, "kind": kind, "object": obj})
        saved.save()
        form.save_m2m()
        _audit(request, "UPDATE", saved)
        messages.success(request, f"{label} updated.")
        return redirect("transport_pro:list", kind=kind)
    return render(request, "transport_pro/object_form.html", {"form": form, "title": "Edit " + label, "kind": kind, "object": obj})


@login_required
@require_POST
def trip_action(request, pk, action):
    if not _has_transport_permission(request.user, "change_transporttrip", "transporttrip"):
        return HttpResponseForbidden("You do not have permission to change trip status.")
    trip = get_object_or_404(models.TransportTrip, pk=pk)
    if action == "authorise":
        if trip.vehicle.documents_expired:
            messages.error(request, "Trip blocked: vehicle documents have expired.")
        elif trip.vehicle.status in ("MAINTENANCE","OUT_OF_SERVICE"):
            messages.error(request, "Trip blocked: vehicle is not serviceable.")
        elif not trip.driver.active or (trip.driver.licence_expiry and trip.driver.licence_expiry < timezone.localdate()):
            messages.error(request, "Trip blocked: driver is inactive or licence expired.")
        elif not hasattr(trip, "inspection") or not trip.inspection.dispatch_allowed:
            messages.error(request, "Trip blocked: a passing pre-trip inspection is required.")
        else:
            trip.status = "AUTHORISED"
            trip.authorised_by = request.user
            trip.save()
            _audit(request, "AUTHORISE_TRIP", trip)
            messages.success(request, "Trip authorised.")
    elif action == "start":
        if trip.status != "AUTHORISED":
            messages.error(request, "Only an authorised trip can start.")
        else:
            trip.status = "IN_PROGRESS"
            trip.actual_departure = timezone.now()
            trip.start_odometer_km = trip.vehicle.odometer_km
            trip.save()
            _audit(request, "START_TRIP", trip)
            messages.success(request, "Trip started.")
    elif action == "complete":
        if trip.status != "IN_PROGRESS":
            messages.error(request, "Only an in-progress trip can be completed.")
        elif trip.boarding_events.filter(status="BOARDED").count() > trip.vehicle.seating_capacity:
            messages.error(request, "Trip cannot close: recorded boarding exceeds vehicle capacity.")
        else:
            trip.status = "COMPLETED"
            trip.actual_arrival = timezone.now()
            end = request.POST.get("end_odometer_km")
            fuel = request.POST.get("fuel_litres")
            if end:
                try:
                    trip.end_odometer_km = int(end)
                    if trip.end_odometer_km < 0 or (trip.start_odometer_km is not None and trip.end_odometer_km < trip.start_odometer_km):
                        raise ValueError()
                except ValueError:
                    messages.error(request, "End odometer must be a valid value not lower than the start odometer.")
                    return redirect("transport_pro:list", kind="trips")
            if fuel:
                try:
                    trip.fuel_litres = Decimal(fuel)
                    if not trip.fuel_litres.is_finite() or trip.fuel_litres < 0:
                        raise ValueError()
                except Exception:
                    messages.error(request, "Fuel litres must be a finite, non-negative number.")
                    return redirect("transport_pro:list", kind="trips")
            trip.closing_notes = request.POST.get("closing_notes","")
            trip.save()
            if trip.end_odometer_km is not None:
                trip.vehicle.odometer_km = max(trip.vehicle.odometer_km, trip.end_odometer_km)
                trip.vehicle.save(update_fields=["odometer_km","updated_at"])
            _audit(request, "COMPLETE_TRIP", trip)
            messages.success(request, "Trip closed. Please ensure every learner is accounted for.")
    else:
        messages.error(request, "Unknown trip action.")
    return redirect("transport_pro:list", kind="trips")


@login_required
def reports(request):
    if not _has_transport_reports_access(request.user):
        return HttpResponseForbidden(
            "You do not have permission to view all transport report summaries."
        )
    today = timezone.localdate()
    by_route = models.LearnerTransport.objects.filter(status="ACTIVE").values("route__code","route__name").annotate(learners=Count("id")).order_by("route__code")
    trips = models.TransportTrip.objects.filter(service_date__year=today.year)
    return render(request, "transport_pro/reports.html", {
        "by_route": by_route, "trips_year": trips.count(),
        "completed_year": trips.filter(status="COMPLETED").count(),
        "incident_count": models.TransportIncident.objects.filter(occurred_at__year=today.year).count(),
        "maintenance_cost": models.MaintenanceRecord.objects.filter(reported_date__year=today.year).aggregate(v=Sum("cost"))["v"] or 0,
        "charges_due": _outstanding_transport_balance(),
    })


@login_required
@require_POST
def verify_payment(request, pk):
    # Superusers may verify. Other users need an active custom ERP role
    # explicitly granted the transport payment change permission.
    if not _can_verify_transport_payments(request.user):
        return HttpResponseForbidden(
            "You do not have permission to verify transport payments."
        )

    with transaction.atomic():
        payment = get_object_or_404(
            models.TransportPayment.objects.select_for_update().select_related("charge"),
            pk=pk,
        )

        if payment.verified:
            messages.info(request, "This payment has already been verified.")
            return redirect("transport_pro:list", kind="payments")

        charge = models.TransportCharge.objects.select_for_update().get(
            pk=payment.charge_id
        )

        payment.verified = True
        payment.verified_by = request.user
        payment.save(
            update_fields=["verified", "verified_by", "updated_at"]
        )

        paid = (
            models.TransportPayment.objects.filter(
                charge=charge, verified=True
            ).aggregate(v=Sum("amount"))["v"]
            or Decimal("0")
        )

        if charge.status != "WAIVED":
            if paid >= charge.amount:
                charge.status = "PAID"
            elif paid > 0:
                charge.status = "PARTIAL"
            else:
                charge.status = "DUE"

            charge.save(update_fields=["status", "updated_at"])

        _audit(
            request,
            "VERIFY_PAYMENT",
            payment,
            {"verified_amount": str(paid)},
        )

    messages.success(
        request,
        "Payment marked verified and charge balance recalculated. "
        "Confirm the transaction against the payment provider separately.",
    )
    return redirect("transport_pro:list", kind="payments")


@login_required
def parent_portal(request):
    """Parent-only transport information limited to their linked children."""
    from decimal import Decimal
    from django.db.models import Sum
    from django.shortcuts import redirect
    from students.models import Student
    from . import models

    parent = getattr(request.user, "parent", None)
    if parent is None:
        return redirect("accounts:login")

    children = list(parent.children.all().order_by("admission_no"))
    child_by_admission = {
        str(child.admission_no).strip(): child
        for child in children
        if child.admission_no
    }
    admission_numbers = list(child_by_admission.keys())

    records = []
    if admission_numbers:
        assignments = (
            models.LearnerTransport.objects
            .filter(learner_admission_no__in=admission_numbers)
            .select_related("route", "stop")
            .order_by("learner_name", "-start_date", "-id")
        )

        for learner in assignments:
            admission = str(learner.learner_admission_no).strip()
            child = child_by_admission.get(admission)
            if child is None:
                continue

            charge_rows = []
            total_charged = Decimal("0.00")
            total_paid = Decimal("0.00")

            for charge in models.TransportCharge.objects.filter(
                learner=learner
            ).order_by("-due_date", "-id"):
                paid = (
                    models.TransportPayment.objects
                    .filter(charge=charge, verified=True)
                    .aggregate(total=Sum("amount"))["total"]
                    or Decimal("0.00")
                )
                amount = charge.amount or Decimal("0.00")
                balance = max(Decimal("0.00"), amount - paid)

                total_charged += amount
                total_paid += paid
                charge_rows.append({
                    "period_label": charge.period_label,
                    "amount": amount,
                    "due_date": charge.due_date,
                    "status": charge.status,
                    "paid": paid,
                    "balance": balance,
                })

            records.append({
                "child": child,
                "learner": learner,
                "charges": charge_rows,
                "total_charged": total_charged,
                "total_paid": total_paid,
                "total_balance": max(
                    Decimal("0.00"), total_charged - total_paid
                ),
            })

    return render(request, "transport_pro/parent_portal.html", {
        "parent": parent,
        "children": children,
        "transport_records": records,
    })
